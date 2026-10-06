"""DOCX rendering for validated final resumes."""

from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.final_resume_models import FinalResume
from app.final_resume_validator import validate_final_resume
from app.templates.registry import get_template


DEFAULT_FONT_NAME = "Arial"
DEFAULT_FONT_SIZE = 10
DEFAULT_HEADING_SIZE = 12
DEFAULT_NAME_SIZE = 18

SECTION_LABELS = {
    "summary": "Professional Summary",
    "skills": "Skills",
    "experience": "Experience",
    "projects": "Projects",
    "education": "Education",
    "certifications": "Certifications",
}


def _safe_text(value):
    """Convert a value into renderable text."""

    if value is None:
        return ""

    if isinstance(value, (list, tuple, set)):
        return ", ".join(
            str(item)
            for item in value
            if item is not None
        )

    if isinstance(value, dict):
        return " | ".join(
            str(item)
            for item in value.values()
            if item is not None
        )

    return str(value)


def _format_paragraph(
    paragraph,
    *,
    space_before=0,
    space_after=0,
    line_spacing=1.0,
    keep_with_next=False,
    keep_together=False,
):
    """Apply consistent resume paragraph layout settings."""

    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_with_next
    fmt.keep_together = keep_together


def _extract_name(personal_info):
    """Return the candidate name."""

    if not isinstance(personal_info, dict):
        return ""

    return _safe_text(
        personal_info.get("name", "")
    )


def _build_contact_line(personal_info):
    """Build a simple contact-information line."""

    if not isinstance(personal_info, dict):
        return ""

    values = []

    for field in (
        "email",
        "phone",
        "location",
        "linkedin",
        "github",
        "website",
    ):
        value = personal_info.get(field)

        if value:
            values.append(
                _safe_text(value)
            )

    return " | ".join(values)


def _extract_title(item):
    """Extract a useful title from a resume item."""

    if not isinstance(item, dict):
        return ""

    for field in (
        "job_title",
        "title",
        "name",
        "degree",
        "company",
    ):
        value = item.get(field)

        if value:
            return _safe_text(value)

    return ""


def _extract_secondary_line(item):
    """Extract secondary information for an item."""

    if not isinstance(item, dict):
        return ""

    values = []

    for field in (
        "company",
        "organization",
        "institution",
        "location",
        "dates",
        "date",
        "year",
        "years",
    ):
        value = item.get(field)

        if value:
            values.append(
                _safe_text(value)
            )

    return " | ".join(values)
def _clean_bullet_text(value):
    """Remove existing bullet markers before Word applies its own bullet."""

    text = _safe_text(value).strip()

    while text.startswith(("•", "", "-", "–", "—")):
        text = text[1:].strip()

    return text

def _extract_bullets(item):
    """Extract bullet-style content without duplicating fallback fields."""

    if isinstance(item, str):
        text = item.strip()
        return [text] if text else []

    if not isinstance(item, dict):
        text = _safe_text(item).strip()
        return [text] if text else []

    preferred_fields = (
        "bullets",
        "responsibilities",
        "achievements",
        "descriptions",
        "description",
    )

    for field in preferred_fields:
        value = item.get(field)

        if value is None:
            continue

        if isinstance(value, (list, tuple)):
            lines = [
                _safe_text(entry).strip()
                for entry in value
                if _safe_text(entry).strip()
            ]

            if lines:
                return lines

        else:
            text = _safe_text(value).strip()

            if text:
                return [text]

    return []


def _set_document_defaults(document, template_config):
    """Apply template-specific document formatting."""

    section = document.sections[0]

    section.top_margin = Inches(
        template_config.top_margin
    )
    section.bottom_margin = Inches(
        template_config.bottom_margin
    )
    section.left_margin = Inches(
        template_config.left_margin
    )
    section.right_margin = Inches(
        template_config.right_margin
    )

    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = template_config.font_name
    normal.font.size = Pt(
        template_config.body_font_size
    )

    # Set East Asia font as well where supported.
    normal._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        template_config.font_name,
    )


def _add_name(document, name, template_config):
    """Add candidate name using template-specific formatting."""

    paragraph = document.add_paragraph()

    alignment = template_config.name_alignment

    if alignment == "left":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif alignment == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(name)
    run.bold = template_config.name_bold
    run.font.name = template_config.font_name
    run.font.size = Pt(
        template_config.name_font_size
    )

    _format_paragraph(
        paragraph,
        space_before=0,
        space_after=2,
        line_spacing=template_config.line_spacing,
        keep_with_next=True,
    )

    return paragraph


def _add_contact_line(
    document,
    contact_line,
    template_config,
):
    """Add contact information using template formatting."""

    if not contact_line:
        return None

    paragraph = document.add_paragraph()

    alignment = template_config.contact_alignment

    if alignment == "left":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif alignment == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(contact_line)
    run.font.name = template_config.font_name
    run.font.size = Pt(
        template_config.contact_font_size
    )

    _format_paragraph(
        paragraph,
        space_before=0,
        space_after=6,
        line_spacing=template_config.line_spacing,
        keep_with_next=True,
    )

    return paragraph


def _add_section_heading(
    document,
    title,
    template_config,
):
    """Add a resume section heading using template formatting."""

    paragraph = document.add_paragraph()

    _format_paragraph(
        paragraph,
        space_before=template_config.section_space_before,
        space_after=template_config.section_space_after,
        line_spacing=template_config.line_spacing,
        keep_with_next=True,
    )

    run = paragraph.add_run(title)
    run.bold = template_config.heading_bold
    run.font.name = template_config.font_name
    run.font.size = Pt(
        template_config.heading_font_size
    )

    return paragraph


def _add_summary(document, summary, template_config):
    """Render professional summary using template formatting."""

    if not summary:
        return

    paragraph = document.add_paragraph()

    run = paragraph.add_run(_safe_text(summary))
    run.font.name = template_config.font_name
    run.font.size = Pt(template_config.body_font_size)

    _format_paragraph(
        paragraph,
        space_before=0,
        space_after=template_config.body_space_after,
        line_spacing=template_config.line_spacing,
        keep_together=True,
    )


def _add_skills(document, skills, template_config):
    """Render skills while preserving optional category information."""

    if not skills:
        return

    if isinstance(skills, dict):
        values = []

        for key, value in skills.items():
            if isinstance(value, (list, tuple, set)):
                items = [
                    _safe_text(item).strip()
                    for item in value
                    if _safe_text(item).strip()
                ]

                if items:
                    values.append(
                        f"{_safe_text(key)}: {', '.join(items)}"
                    )

            elif value:
                values.append(
                    f"{_safe_text(key)}: {_safe_text(value)}"
                )

    elif isinstance(skills, (list, tuple, set)):
        values = [
            _safe_text(item).strip()
            for item in skills
            if _safe_text(item).strip()
        ]

    else:
        values = [_safe_text(skills).strip()]

    values = [
        value
        for value in values
        if value
    ]

    if not values:
        return

    paragraph = document.add_paragraph()

    run = paragraph.add_run(", ".join(values))
    run.font.name = template_config.font_name
    run.font.size = Pt(template_config.body_font_size)

    _format_paragraph(
        paragraph,
        space_before=0,
        space_after=template_config.body_space_after,
        line_spacing=template_config.line_spacing,
        keep_together=True,
    )

    
def _add_item(document, item, template_config):
    """Render one experience/project/education item."""

    title = _extract_title(item)

    if title:
        paragraph = document.add_paragraph()

        _format_paragraph(
            paragraph,
            space_before=template_config.section_space_after,
            space_after=0,
            line_spacing=template_config.line_spacing,
            keep_with_next=True,
            keep_together=True,
        )

        run = paragraph.add_run(title)
        run.bold = template_config.heading_bold
        run.font.name = template_config.font_name
        run.font.size = Pt(template_config.item_title_font_size)

    secondary = _extract_secondary_line(item)

    if secondary:
        paragraph = document.add_paragraph()

        _format_paragraph(
            paragraph,
            space_before=0,
            space_after=template_config.body_space_after,
            line_spacing=template_config.line_spacing,
            keep_with_next=True,
            keep_together=True,
        )

        run = paragraph.add_run(secondary)
        run.italic = True
        run.font.name = template_config.font_name
        run.font.size = Pt(template_config.secondary_font_size)

    bullets = _extract_bullets(item)

    for bullet in bullets:
        if not template_config.use_bullets:
            paragraph = document.add_paragraph()
        else:
            paragraph = document.add_paragraph(
                style="List Bullet"
            )

        _format_paragraph(
            paragraph,
            space_before=0,
            space_after=template_config.body_space_after,
            line_spacing=template_config.line_spacing,
            keep_together=True,
        )

        cleaned_bullet = _clean_bullet_text(
            bullet
        )

        if not cleaned_bullet:
            continue

        run = paragraph.add_run(
            cleaned_bullet
        )

        run.font.name = template_config.font_name
        run.font.size = Pt(template_config.body_font_size)


def _add_collection(
    document,
    items,
    template_config,
):
    """Render a collection of resume items."""

    for item in items or []:
        _add_item(
            document,
            item,
            template_config,
        )


def render_final_resume(
    resume,
    output_path,
    validate=True,
    template="default",
):
    """
    Render a FinalResume to DOCX.

    Args:
        resume: FinalResume instance.
        output_path: destination .docx path.
        validate: validate before rendering.

    Returns:
        Path to generated DOCX.
    """

    if not isinstance(
        resume,
        FinalResume,
    ):
        raise TypeError(
            "resume must be a FinalResume instance."
        )

    template_config = get_template(template)

    if validate:
        validation = validate_final_resume(
            resume
        )

        if not validation["valid"]:
            raise ValueError(
                "Cannot render invalid resume: "
                + "; ".join(
                    validation["errors"]
                )
            )

    output_path = Path(output_path)

    if output_path.suffix.lower() != ".docx":
        raise ValueError(
            "output_path must have a .docx extension."
        )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # Rendering should never mutate the FinalResume.
    resume = deepcopy(resume)

    document = Document()

    _set_document_defaults(
        document,
        template_config,
    )

    personal_info = resume.personal_info

    name = _extract_name(personal_info)

    if name:
        _add_name(
        document,
        name,
        template_config,
    )

    contact_line = _build_contact_line(
        personal_info
    )

    _add_contact_line(
        document,
        contact_line,
        template_config,
    )

    section_values = {
        "summary": resume.summary,
        "skills": resume.skills,
        "experience": resume.experience,
        "projects": resume.projects,
        "education": resume.education,
        "certifications": resume.certifications,
    }

    for section in resume.section_order:
        value = section_values.get(
            section
        )

        if not value:
            continue

        heading = SECTION_LABELS.get(
            section,
            section.title(),
        )

        _add_section_heading(
            document,
            heading,
            template_config,
        )

        if section == "summary":
            _add_summary(
                document,
                value,
                template_config,
            )
        elif section == "skills":
            _add_skills(
                document,
                value,
                template_config,
            )
        else:
            _add_collection(
                document,
                value,
                template_config,
            )

    document.save(
        str(output_path)
    )

    return output_path
