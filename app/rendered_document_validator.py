"""Validation helpers for rendered resume documents."""

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx"}


def count_pdf_pages(pdf_path):
    """Return the actual number of pages in a rendered PDF."""
    path = Path(pdf_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError("pdf_path must use the .pdf extension.")

    if not path.exists():
        raise FileNotFoundError(
            f"PDF file does not exist: {path}"
        )

    try:
        reader = PdfReader(str(path))
        return len(reader.pages)
    except Exception as exc:
        raise ValueError(
            f"Could not read rendered PDF: {path}"
        ) from exc


def validate_pdf_page_count(pdf_path, target_pages):
    """Compare actual PDF pages with the requested target page count."""
    if (
        not isinstance(target_pages, int)
        or isinstance(target_pages, bool)
        or target_pages < 1
    ):
        raise ValueError(
            "target_pages must be a positive integer."
        )

    actual_pages = count_pdf_pages(pdf_path)

    return {
        "valid": actual_pages == target_pages,
        "target_pages": target_pages,
        "actual_pages": actual_pages,
        "page_count_match": actual_pages == target_pages,
        "message": (
            "PDF page count matches target."
            if actual_pages == target_pages
            else (
                "PDF page count does not match target: "
                f"expected {target_pages}, got {actual_pages}."
            )
        ),
    }


def extract_pdf_text(pdf_path):
    """Extract text from all pages of a rendered PDF."""
    path = Path(pdf_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError("pdf_path must use the .pdf extension.")

    if not path.exists():
        raise FileNotFoundError(
            f"PDF file does not exist: {path}"
        )

    try:
        reader = PdfReader(str(path))

        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text)

        return "\n".join(pages).strip()

    except Exception as exc:
        raise ValueError(
            f"Could not extract text from rendered PDF: {path}"
        ) from exc


def extract_docx_text(docx_path):
    """Extract paragraph text from a rendered DOCX."""
    path = Path(docx_path)

    if path.suffix.lower() != ".docx":
        raise ValueError("docx_path must use the .docx extension.")

    if not path.exists():
        raise FileNotFoundError(
            f"DOCX file does not exist: {path}"
        )

    try:
        document = Document(str(path))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return "\n".join(paragraphs).strip()

    except Exception as exc:
        raise ValueError(
            f"Could not read rendered DOCX: {path}"
        ) from exc


def validate_pdf_structure(pdf_path):
    """Validate the basic structural integrity of a rendered PDF."""
    path = Path(pdf_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError("pdf_path must use the .pdf extension.")

    if not path.exists():
        raise FileNotFoundError(
            f"PDF file does not exist: {path}"
        )

    try:
        reader = PdfReader(str(path))
        page_count = len(reader.pages)

        text = extract_pdf_text(path)

        return {
            "valid": page_count > 0 and bool(text.strip()),
            "file_exists": True,
            "readable": True,
            "page_count": page_count,
            "text_present": bool(text.strip()),
            "message": (
                "PDF structure is valid."
                if page_count > 0 and text.strip()
                else "PDF does not contain readable resume text."
            ),
        }

    except Exception as exc:
        return {
            "valid": False,
            "file_exists": path.exists(),
            "readable": False,
            "page_count": 0,
            "text_present": False,
            "message": f"PDF structural validation failed: {exc}",
        }

def validate_final_pdf(
    pdf_path,
    required_content=None,
    target_pages=None,
):
    """Run final ATS-oriented validation on a rendered PDF.

    Parameters
    ----------
    pdf_path:
        Path to the rendered PDF file.

    required_content:
        Optional iterable of strings that must be present in the
        extracted PDF text.

    target_pages:
        Optional expected PDF page count.

    Returns
    -------
    dict
        Structured final PDF validation result.
    """

    path = Path(pdf_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError(
            "pdf_path must use the .pdf extension."
        )

    if not path.exists():
        return {
            "valid": False,
            "file_exists": False,
            "readable": False,
            "text_present": False,
            "required_content_present": False,
            "page_count_valid": False,
            "all_pages_have_text": False,
            "standard_headings_present": False,
            "missing_content": tuple(
                required_content or ()
            ),
            "message": (
                f"PDF file does not exist: {path}"
            ),
        }

    try:
        reader = PdfReader(str(path))

        page_count = len(reader.pages)

        text = extract_pdf_text(path)

        required_content = tuple(
            str(value).strip()
            for value in (required_content or ())
            if str(value).strip()
        )

        missing_content = tuple(
            value
            for value in required_content
            if value not in text
        )

        required_content_present = (
            len(missing_content) == 0
        )

        page_texts = [
            (page.extract_text() or "").strip()
            for page in reader.pages
        ]

        all_pages_have_text = (
            page_count > 0
            and all(page_texts)
        )

        if target_pages is None:
            page_count_valid = page_count > 0
        else:
            if (
                not isinstance(target_pages, int)
                or isinstance(target_pages, bool)
                or target_pages < 1
            ):
                raise ValueError(
                    "target_pages must be a positive integer."
                )

            page_count_valid = (
                page_count == target_pages
            )

        text_upper = text.upper()

        standard_headings = (
            "SUMMARY",
            "SKILLS",
            "EXPERIENCE",
            "PROJECTS",
            "EDUCATION",
            "CERTIFICATIONS",
        )

        present_headings = tuple(
            heading
            for heading in standard_headings
            if heading in text_upper
        )

        standard_headings_present = (
            len(present_headings)
            == len(standard_headings)
        )

        text_present = bool(
            text.strip()
        )

        valid = all(
            (
                text_present,
                required_content_present,
                page_count_valid,
                all_pages_have_text,
                standard_headings_present,
            )
        )

        failures = []

        if not text_present:
            failures.append(
                "PDF contains no extractable text."
            )

        if not required_content_present:
            failures.append(
                "Required resume content is missing."
            )

        if not page_count_valid:
            if target_pages is None:
                failures.append(
                    "PDF does not contain any pages."
                )
            else:
                failures.append(
                    "PDF page count does not match target."
                )

        if not all_pages_have_text:
            failures.append(
                "One or more PDF pages contain no "
                "extractable text."
            )

        if not standard_headings_present:
            failures.append(
                "PDF is missing one or more standard "
                "resume section headings."
            )

        return {
            "valid": valid,
            "file_exists": True,
            "readable": True,
            "text_present": text_present,
            "required_content_present": (
                required_content_present
            ),
            "page_count": page_count,
            "target_pages": target_pages,
            "page_count_valid": page_count_valid,
            "all_pages_have_text": (
                all_pages_have_text
            ),
            "standard_headings_present": (
                standard_headings_present
            ),
            "present_headings": present_headings,
            "missing_content": missing_content,
            "message": (
                "Final PDF validation passed."
                if valid
                else (
                    "Final PDF validation failed: "
                    + " ".join(failures)
                )
            ),
        }

    except ValueError:
        raise

    except Exception as exc:
        return {
            "valid": False,
            "file_exists": path.exists(),
            "readable": False,
            "text_present": False,
            "required_content_present": False,
            "page_count": 0,
            "target_pages": target_pages,
            "page_count_valid": False,
            "all_pages_have_text": False,
            "standard_headings_present": False,
            "missing_content": tuple(
                required_content or ()
            ),
            "message": (
                f"Final PDF validation failed: {exc}"
            ),
        }
    
def validate_docx_structure(docx_path):
    """Validate the basic structural integrity of a rendered DOCX."""
    path = Path(docx_path)

    if path.suffix.lower() != ".docx":
        raise ValueError("docx_path must use the .docx extension.")

    if not path.exists():
        raise FileNotFoundError(
            f"DOCX file does not exist: {path}"
        )

    try:
        document = Document(str(path))

        paragraph_count = len(document.paragraphs)
        text = extract_docx_text(path)

        return {
            "valid": paragraph_count > 0 and bool(text.strip()),
            "file_exists": True,
            "readable": True,
            "paragraph_count": paragraph_count,
            "text_present": bool(text.strip()),
            "message": (
                "DOCX structure is valid."
                if paragraph_count > 0 and text.strip()
                else "DOCX does not contain readable resume text."
            ),
        }

    except Exception as exc:
        return {
            "valid": False,
            "file_exists": path.exists(),
            "readable": False,
            "paragraph_count": 0,
            "text_present": False,
            "message": f"DOCX structural validation failed: {exc}",
        }

def validate_final_docx(
    docx_path,
    required_content=None,
):
    """Run final ATS-oriented validation on a rendered DOCX.

    This extends the existing structural DOCX validation with checks
    required for the final Phase 15 validation gate.

    Parameters
    ----------
    docx_path:
        Path to the rendered DOCX file.

    required_content:
        Optional iterable of strings that must be present in the
        extracted document text.

    Returns
    -------
    dict
        Structured final DOCX validation result.
    """

    path = Path(docx_path)

    if path.suffix.lower() != ".docx":
        raise ValueError(
            "docx_path must use the .docx extension."
        )

    if not path.exists():
        return {
            "valid": False,
            "file_exists": False,
            "readable": False,
            "text_present": False,
            "required_content_present": False,
            "no_tables": False,
            "no_textboxes": False,
            "standard_headings_present": False,
            "missing_content": tuple(
                required_content or ()
            ),
            "message": (
                f"DOCX file does not exist: {path}"
            ),
        }

    try:
        document = Document(str(path))

        text = extract_docx_text(path)

        # -----------------------------------------------------------
        # Required content
        # -----------------------------------------------------------

        required_content = tuple(
            str(value).strip()
            for value in (required_content or ())
            if str(value).strip()
        )

        missing_content = tuple(
            value
            for value in required_content
            if value not in text
        )

        required_content_present = (
            len(missing_content) == 0
        )

        # -----------------------------------------------------------
        # ATS structural checks
        # -----------------------------------------------------------

        no_tables = (
            len(document.tables) == 0
        )

        # Inspect the underlying DOCX XML for textboxes/shapes.
        with ZipFile(path) as archive:
            document_xml = archive.read(
                "word/document.xml"
            ).decode(
                "utf-8",
                errors="ignore",
            )

        no_textboxes = (
            "<w:txbxContent" not in document_xml
            and "<wps:" not in document_xml
            and "<v:shape" not in document_xml
        )

        # -----------------------------------------------------------
        # Standard resume headings
        # -----------------------------------------------------------

        text_upper = text.upper()

        standard_headings = (
            "SUMMARY",
            "SKILLS",
            "EXPERIENCE",
            "PROJECTS",
            "EDUCATION",
            "CERTIFICATIONS",
        )

        present_headings = tuple(
            heading
            for heading in standard_headings
            if heading in text_upper
        )

        standard_headings_present = (
            len(present_headings)
            == len(standard_headings)
        )

        # -----------------------------------------------------------
        # Overall validation
        # -----------------------------------------------------------

        text_present = bool(
            text.strip()
        )

        valid = all(
            (
                text_present,
                required_content_present,
                no_tables,
                no_textboxes,
                standard_headings_present,
            )
        )

        failures = []

        if not text_present:
            failures.append(
                "DOCX contains no extractable text."
            )

        if not required_content_present:
            failures.append(
                "Required resume content is missing."
            )

        if not no_tables:
            failures.append(
                "DOCX contains tables."
            )

        if not no_textboxes:
            failures.append(
                "DOCX contains textboxes or shapes."
            )

        if not standard_headings_present:
            failures.append(
                "DOCX is missing one or more standard "
                "resume section headings."
            )

        return {
            "valid": valid,
            "file_exists": True,
            "readable": True,
            "text_present": text_present,
            "required_content_present": (
                required_content_present
            ),
            "no_tables": no_tables,
            "no_textboxes": no_textboxes,
            "standard_headings_present": (
                standard_headings_present
            ),
            "present_headings": present_headings,
            "missing_content": missing_content,
            "message": (
                "Final DOCX validation passed."
                if valid
                else (
                    "Final DOCX validation failed: "
                    + " ".join(failures)
                )
            ),
        }

    except Exception as exc:
        return {
            "valid": False,
            "file_exists": path.exists(),
            "readable": False,
            "text_present": False,
            "required_content_present": False,
            "no_tables": False,
            "no_textboxes": False,
            "standard_headings_present": False,
            "missing_content": tuple(
                required_content or ()
            ),
            "message": (
                f"Final DOCX validation failed: {exc}"
            ),
        }
    
def validate_rendered_documents(
    docx_path,
    pdf_path,
    target_pages,
):
    """Validate both rendered resume documents."""
    pdf_structure = validate_pdf_structure(pdf_path)
    docx_structure = validate_docx_structure(docx_path)
    page_count = validate_pdf_page_count(
        pdf_path,
        target_pages,
    )

    overall_valid = (
        pdf_structure["valid"]
        and docx_structure["valid"]
        and page_count["valid"]
    )

    return {
        "valid": overall_valid,
        "pdf": {
            "valid": pdf_structure["valid"],
            "file_exists": pdf_structure["file_exists"],
            "readable": pdf_structure["readable"],
            "page_count": page_count,
            "text_present": pdf_structure["text_present"],
            "message": pdf_structure["message"],
        },
        "docx": docx_structure,
        "message": (
            "Rendered documents are structurally valid."
            if overall_valid
            else "One or more rendered document validations failed."
        ),
    }

def validate_final_documents(
    docx_path,
    pdf_path,
    ats_comparison,
    resume_text,
    required_content=None,
    target_pages=None,
):
    """Run the complete final validation for generated documents.

    Combines ATS safety validation with final DOCX and PDF
    document validation.
    """

    from app.final_ats_validation import run_ats_safety_gate

    ats_safety = run_ats_safety_gate(
        ats_comparison,
        resume_text,
    )

    docx_validation = validate_final_docx(
        docx_path,
        required_content=required_content,
    )

    pdf_validation = validate_final_pdf(
        pdf_path,
        required_content=required_content,
        target_pages=target_pages,
    )

    valid = all(
        (
            ats_safety.passed,
            docx_validation["valid"],
            pdf_validation["valid"],
        )
    )

    failures = []

    if not ats_safety.passed:
        failures.append(
            "ATS safety gate failed."
        )

    if not docx_validation["valid"]:
        failures.append(
            "Final DOCX validation failed."
        )

    if not pdf_validation["valid"]:
        failures.append(
            "Final PDF validation failed."
        )

    return {
        "valid": valid,
        "ats_safety": ats_safety,
        "docx": docx_validation,
        "pdf": pdf_validation,
        "message": (
            "Final resume document validation passed."
            if valid
            else (
                "Final resume document validation failed: "
                + " ".join(failures)
            )
        ),
    }