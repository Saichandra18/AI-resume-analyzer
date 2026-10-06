from pathlib import Path

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.final_ats_validation import (
    compare_ats_results,
)

from app.rendered_document_validator import (
    count_pdf_pages,
    extract_docx_text,
    extract_pdf_text,
    validate_docx_structure,
    validate_final_docx,
    validate_final_pdf,
    validate_final_documents,
    validate_pdf_page_count,
    validate_pdf_structure,
    validate_rendered_documents,
)


def create_pdf(path, pages=1, text="Test Resume"):
    """Create a valid PDF with the requested number of pages."""
    pdf = canvas.Canvas(str(path))

    for page_number in range(pages):
        pdf.drawString(
            72,
            720,
            f"{text} - Page {page_number + 1}",
        )
        pdf.showPage()

    pdf.save()


def create_docx(path, text="Test Resume"):
    """Create a valid DOCX containing resume text."""
    document = Document()
    document.add_paragraph(text)
    document.save(path)


def test_count_pdf_pages_returns_one_for_one_page_pdf(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(pdf_path, pages=1)

    assert count_pdf_pages(pdf_path) == 1


def test_count_pdf_pages_returns_two_for_two_page_pdf(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(pdf_path, pages=2)

    assert count_pdf_pages(pdf_path) == 2


def test_count_pdf_pages_rejects_non_pdf_extension(tmp_path):
    path = tmp_path / "resume.docx"

    path.write_bytes(b"test")

    with pytest.raises(
        ValueError,
        match=r"pdf_path must use the \.pdf extension",
    ):
        count_pdf_pages(path)


def test_count_pdf_pages_rejects_missing_file(tmp_path):
    pdf_path = tmp_path / "missing.pdf"

    with pytest.raises(FileNotFoundError):
        count_pdf_pages(pdf_path)


def test_validate_pdf_page_count_accepts_matching_count(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(pdf_path, pages=1)

    result = validate_pdf_page_count(
        pdf_path,
        target_pages=1,
    )

    assert result["valid"] is True
    assert result["target_pages"] == 1
    assert result["actual_pages"] == 1
    assert result["page_count_match"] is True


def test_validate_pdf_page_count_detects_mismatch(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(pdf_path, pages=2)

    result = validate_pdf_page_count(
        pdf_path,
        target_pages=1,
    )

    assert result["valid"] is False
    assert result["target_pages"] == 1
    assert result["actual_pages"] == 2
    assert result["page_count_match"] is False


def test_validate_pdf_page_count_rejects_invalid_target_pages(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(pdf_path, pages=1)

    with pytest.raises(
        ValueError,
        match="target_pages must be a positive integer",
    ):
        validate_pdf_page_count(
            pdf_path,
            target_pages=0,
        )


def test_extract_pdf_text_returns_rendered_text(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(
        pdf_path,
        pages=1,
        text="AI ML Engineer",
    )

    text = extract_pdf_text(pdf_path)

    assert "AI ML Engineer" in text


def test_extract_docx_text_returns_rendered_text(tmp_path):
    docx_path = tmp_path / "resume.docx"

    create_docx(
        docx_path,
        text="AI ML Engineer",
    )

    text = extract_docx_text(docx_path)

    assert "AI ML Engineer" in text


def test_validate_pdf_structure_accepts_valid_pdf(tmp_path):
    pdf_path = tmp_path / "resume.pdf"

    create_pdf(
        pdf_path,
        pages=1,
        text="AI ML Engineer",
    )

    result = validate_pdf_structure(pdf_path)

    assert result["valid"] is True
    assert result["file_exists"] is True
    assert result["readable"] is True
    assert result["page_count"] == 1
    assert result["text_present"] is True


def test_validate_docx_structure_accepts_valid_docx(tmp_path):
    docx_path = tmp_path / "resume.docx"

    create_docx(
        docx_path,
        text="AI ML Engineer",
    )

    result = validate_docx_structure(docx_path)

    assert result["valid"] is True
    assert result["file_exists"] is True
    assert result["readable"] is True
    assert result["paragraph_count"] >= 1
    assert result["text_present"] is True


def test_validate_rendered_documents_accepts_valid_documents(
    tmp_path,
):
    docx_path = tmp_path / "resume.docx"
    pdf_path = tmp_path / "resume.pdf"

    create_docx(
        docx_path,
        text="AI ML Engineer",
    )

    create_pdf(
        pdf_path,
        pages=1,
        text="AI ML Engineer",
    )

    result = validate_rendered_documents(
        docx_path=docx_path,
        pdf_path=pdf_path,
        target_pages=1,
    )

    assert result["valid"] is True

    assert result["pdf"]["valid"] is True
    assert result["pdf"]["readable"] is True
    assert result["pdf"]["text_present"] is True
    assert result["pdf"]["page_count"]["actual_pages"] == 1

    assert result["docx"]["valid"] is True
    assert result["docx"]["readable"] is True
    assert result["docx"]["text_present"] is True

def test_validate_final_docx_accepts_valid_resume(tmp_path):
    docx_path = tmp_path / "resume.docx"

    document = Document()

    for heading in (
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
    ):
        document.add_paragraph(heading)

    document.add_paragraph(
        "AI/ML Engineer"
    )

    document.add_paragraph(
        "Python Machine Learning SQL"
    )

    document.save(docx_path)

    result = validate_final_docx(
        docx_path,
        required_content=[
            "AI/ML Engineer",
            "Python",
            "Machine Learning",
            "SQL",
        ],
    )

    assert result["valid"] is True, result
    assert result["file_exists"] is True
    assert result["readable"] is True
    assert result["text_present"] is True
    assert result["required_content_present"] is True
    assert result["no_tables"] is True
    assert result["no_textboxes"] is True
    assert result["standard_headings_present"] is True
    assert result["missing_content"] == ()


def test_validate_final_docx_detects_missing_content(
    tmp_path,
):
    docx_path = tmp_path / "resume.docx"

    create_docx(
        docx_path,
        text="AI/ML Engineer",
    )

    result = validate_final_docx(
        docx_path,
        required_content=[
            "Python",
        ],
    )

    assert result["valid"] is False
    assert result["required_content_present"] is False
    assert "Python" in result["missing_content"]


def test_validate_final_docx_rejects_tables(
    tmp_path,
):
    docx_path = tmp_path / "resume.docx"

    document = Document()

    for heading in (
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
    ):
        document.add_paragraph(heading)

    table = document.add_table(
        rows=1,
        cols=1,
    )

    table.cell(0, 0).text = "Python"

    document.save(docx_path)

    result = validate_final_docx(
        docx_path,
    )

    assert result["valid"] is False
    assert result["no_tables"] is False


def test_validate_final_docx_missing_file(
    tmp_path,
):
    docx_path = tmp_path / "missing.docx"

    result = validate_final_docx(
        docx_path,
    )

    assert result["valid"] is False
    assert result["file_exists"] is False


def test_validate_final_docx_rejects_non_docx(
    tmp_path,
):
    path = tmp_path / "resume.pdf"

    path.write_bytes(b"test")

    with pytest.raises(
        ValueError,
        match="docx_path must use the .docx extension",
    ):
        validate_final_docx(path)

def test_validate_final_pdf_accepts_valid_resume(
    tmp_path,
):
    pdf_path = tmp_path / "resume.pdf"

    pdf = canvas.Canvas(str(pdf_path))

    for line in (
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
        "AI/ML Engineer",
        "Python Machine Learning SQL",
    ):
        pdf.drawString(
            72,
            750 - (
                20 * (
                    [
                        "SUMMARY",
                        "SKILLS",
                        "EXPERIENCE",
                        "PROJECTS",
                        "EDUCATION",
                        "CERTIFICATIONS",
                        "AI/ML Engineer",
                        "Python Machine Learning SQL",
                    ].index(line)
                )
            ),
            line,
        )

    pdf.save()

    result = validate_final_pdf(
        pdf_path,
        required_content=[
            "AI/ML Engineer",
            "Python",
            "Machine Learning",
            "SQL",
        ],
        target_pages=1,
    )

    assert result["valid"] is True
    assert result["file_exists"] is True
    assert result["readable"] is True
    assert result["text_present"] is True
    assert result["required_content_present"] is True
    assert result["page_count"] == 1
    assert result["page_count_valid"] is True
    assert result["all_pages_have_text"] is True
    assert result["standard_headings_present"] is True
    assert result["missing_content"] == ()


def test_validate_final_pdf_detects_missing_content(
    tmp_path,
):
    pdf_path = tmp_path / "resume.pdf"

    pdf = canvas.Canvas(str(pdf_path))

    for index, line in enumerate(
        (
            "SUMMARY",
            "SKILLS",
            "EXPERIENCE",
            "PROJECTS",
            "EDUCATION",
            "CERTIFICATIONS",
            "AI/ML Engineer",
        )
    ):
        pdf.drawString(
            72,
            750 - (20 * index),
            line,
        )

    pdf.save()

    result = validate_final_pdf(
        pdf_path,
        required_content=[
            "Python",
        ],
        target_pages=1,
    )

    assert result["valid"] is False
    assert result["required_content_present"] is False
    assert "Python" in result["missing_content"]


def test_validate_final_pdf_detects_page_count_mismatch(
    tmp_path,
):
    pdf_path = tmp_path / "resume.pdf"

    pdf = canvas.Canvas(str(pdf_path))

    pdf.drawString(
        72,
        750,
        "SUMMARY",
    )

    pdf.showPage()

    pdf.drawString(
        72,
        750,
        "SKILLS",
    )

    pdf.save()

    result = validate_final_pdf(
        pdf_path,
        target_pages=1,
    )

    assert result["valid"] is False
    assert result["page_count"] == 2
    assert result["page_count_valid"] is False


def test_validate_final_pdf_missing_file(
    tmp_path,
):
    pdf_path = tmp_path / "missing.pdf"

    result = validate_final_pdf(
        pdf_path,
    )

    assert result["valid"] is False
    assert result["file_exists"] is False


def test_validate_final_pdf_rejects_non_pdf(
    tmp_path,
):
    path = tmp_path / "resume.docx"

    path.write_bytes(b"test")

    with pytest.raises(
        ValueError,
        match="pdf_path must use the .pdf extension",
    ):
        validate_final_pdf(path)

def test_validate_final_documents_accepts_safe_documents(
    tmp_path,
):
    docx_path = tmp_path / "resume.docx"
    pdf_path = tmp_path / "resume.pdf"

    headings = (
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
    )

    # Create DOCX
    document = Document()

    for heading in headings:
        document.add_paragraph(heading)

    document.add_paragraph(
        "AI/ML Engineer"
    )

    document.add_paragraph(
        "Python Machine Learning SQL"
    )

    document.save(docx_path)

    # Create PDF
    pdf = canvas.Canvas(str(pdf_path))

    lines = (
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
        "AI/ML Engineer",
        "Python Machine Learning SQL",
    )

    for index, line in enumerate(lines):
        pdf.drawString(
            72,
            750 - (20 * index),
            line,
        )

    pdf.save()

    ats_comparison = compare_ats_results(
        {
            "ats_score": 70,
            "skills": {
                "matched": ["Python"],
                "missing": [],
            },
            "keywords": {
                "matched": ["Machine Learning"],
                "missing": [],
            },
            "component_scores": {
                "skills": 80,
                "experience": 70,
                "qualifications": 75,
                "keywords": 80,
            },
        },
        {
            "ats_score": 85,
            "skills": {
                "matched": [
                    "Python",
                    "SQL",
                ],
                "missing": [],
            },
            "keywords": {
                "matched": [
                    "Machine Learning",
                    "AI",
                ],
                "missing": [],
            },
            "component_scores": {
                "skills": 90,
                "experience": 75,
                "qualifications": 80,
                "keywords": 90,
            },
        },
    )

    result = validate_final_documents(
        docx_path,
        pdf_path,
        ats_comparison,
        resume_text=(
            "SUMMARY SKILLS EXPERIENCE PROJECTS "
            "EDUCATION CERTIFICATIONS "
            "AI/ML Engineer Python Machine Learning SQL"
        ),
        required_content=[
            "AI/ML Engineer",
            "Python",
            "Machine Learning",
            "SQL",
        ],
        target_pages=1,
    )

    assert result["valid"] is True
    assert result["ats_safety"].passed is True
    assert result["docx"]["valid"] is True
    assert result["pdf"]["valid"] is True