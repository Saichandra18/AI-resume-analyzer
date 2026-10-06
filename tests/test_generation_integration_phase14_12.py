from pathlib import Path

import pytest

from app.resume_generation_service import (
    generate_optimized_resume_docx,
    generate_optimized_resume_pdf,
    generate_optimized_resume_documents,
)

TEMPLATES = ["default", "ats_clean", "modern", "technical"]

SAMPLE_RESUME_TEXT = """
John Doe
john.doe@example.com | +91 9876543210 | Bengaluru, India

PROFESSIONAL SUMMARY
AI/ML Engineer with experience building machine learning and generative AI solutions.

SKILLS
Python, SQL, Machine Learning, Generative AI, RAG, Flask, Pandas

EXPERIENCE
AI/ML Engineer
Tech Company
Developed machine learning and RAG-based applications using Python and SQL.

PROJECTS
AI Resume Analyzer
Built an AI-powered resume analysis and optimization platform.

EDUCATION
Bachelor of Technology in Computer Science
University of Technology
"""

SAMPLE_JOB_DESCRIPTION = """
We are looking for an AI/ML Engineer with skills in Python, SQL,
Machine Learning, Generative AI, RAG, and Flask.
"""

SAMPLE_RESUME_DATA = {
    "personal_info": {
        "name": "John Doe",
        "email": "john.doe@example.com",
        "phone": "+91 9876543210",
        "location": "Bengaluru, India",
    },
    "summary": (
        "AI/ML Engineer with experience building machine learning "
        "and generative AI solutions."
    ),
    "skills": ["Python", "SQL", "Machine Learning", "Generative AI", "RAG", "Flask", "Pandas"],
    "experience": {
        "section": ["Data Analyst", "ABC Technologies", "2021 - Present"],
        "job_titles": ["Data Analyst"],
        "years": ["2021 - Present"],
    },
    "projects": [
        "AI Resume Analyzer",
        "Built an AI-powered resume analysis and optimization platform.",
    ],
    "education": {
        "degrees": ["Bachelor of Technology in Computer Science"],
        "years": ["2020", "2024"],
    },
    "certifications": [],
}

@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_generation_succeeds_for_all_templates(tmp_path, template):
    output_path = tmp_path / f"resume_{template}.docx"
    result = generate_optimized_resume_docx(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template=template)
    assert output_path.exists()
    assert output_path.stat().st_size > 0
    assert result is not None

@pytest.mark.parametrize("template", TEMPLATES)
def test_pdf_generation_succeeds_for_all_templates(tmp_path, template):
    output_path = tmp_path / f"resume_{template}.pdf"
    result = generate_optimized_resume_pdf(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template=template)
    assert output_path.exists()
    assert output_path.stat().st_size > 0
    assert result is not None

@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_generation_preserves_resume_content(tmp_path, template):
    output_path = tmp_path / f"resume_{template}.docx"
    generate_optimized_resume_docx(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template=template)
    assert output_path.exists()
    from docx import Document
    document = Document(str(output_path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "John Doe" in text
    assert "AI/ML Engineer" in text
    assert "Python" in text
    normalized_text = " ".join(text.lower().split())
    assert "resume" in normalized_text
    assert (
        "analyze" in normalized_text
        or "analyz" in normalized_text
        or "analysis" in normalized_text
    )
    assert (
        "optimize" in normalized_text
        or "optimization" in normalized_text
        or "optimizing" in normalized_text
    )

@pytest.mark.parametrize("template", TEMPLATES)
def test_pdf_generation_preserves_resume_content(tmp_path, template):
    output_path = tmp_path / f"resume_{template}.pdf"
    generate_optimized_resume_pdf(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template=template)
    assert output_path.exists()
    from pypdf import PdfReader
    reader = PdfReader(str(output_path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "JOHN DOE" in text.upper()
    assert "AI/ML ENGINEER" in text.upper()
    assert "PYTHON" in text.upper()
    normalized_text = " ".join(text.upper().split())
    assert "RESUME" in normalized_text
    assert "ANALYZ" in normalized_text or "ANALYSIS" in normalized_text
    assert "OPTIMIZ" in normalized_text

@pytest.mark.parametrize("template", TEMPLATES)
def test_combined_generation_creates_both_documents(tmp_path, template):
    docx_path = tmp_path / f"resume_{template}.docx"
    pdf_path = tmp_path / f"resume_{template}.pdf"
    result = generate_optimized_resume_documents(resume_text=SAMPLE_RESUME_TEXT, docx_output_path=str(docx_path), pdf_output_path=str(pdf_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template=template)
    assert docx_path.exists()
    assert pdf_path.exists()
    assert docx_path.stat().st_size > 0
    assert pdf_path.stat().st_size > 0
    assert result is not None

def test_default_template_remains_backward_compatible(tmp_path):
    output_path = tmp_path / "default_resume.docx"
    generate_optimized_resume_docx(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION)
    assert output_path.exists()
    assert output_path.stat().st_size > 0

def test_template_selection_does_not_change_final_resume_content(tmp_path):
    outputs = {}
    for template in TEMPLATES:
        output_path = tmp_path / f"{template}.docx"
        generate_optimized_resume_docx(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template=template)
        from docx import Document
        document = Document(str(output_path))
        outputs[template] = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for template, text in outputs.items():
        assert "John Doe" in text
        assert "AI/ML Engineer" in text
        normalized_text = " ".join(text.lower().split())
        assert "resume" in normalized_text
        assert "analy" in normalized_text
        assert "optim" in normalized_text

def test_invalid_template_is_rejected(tmp_path):
    output_path = tmp_path / "invalid.docx"
    with pytest.raises((KeyError, ValueError)):
        generate_optimized_resume_docx(resume_text=SAMPLE_RESUME_TEXT, output_path=str(output_path), resume_data=SAMPLE_RESUME_DATA, job_description=SAMPLE_JOB_DESCRIPTION, template="does_not_exist")
