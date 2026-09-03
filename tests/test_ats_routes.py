from io import BytesIO
from urllib import response

from app import create_app



def test_ats_analysis_success(tmp_path):
    from docx import Document

    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    file_path = tmp_path / "resume.docx"

    document = Document()

    document.add_paragraph("Sai Chandra")
    document.add_paragraph("sai@example.com")
    document.add_paragraph("+91 9876543210")

    document.add_paragraph("SKILLS")
    document.add_paragraph("Python")
    document.add_paragraph("SQL")
    document.add_paragraph("Pandas")
    document.add_paragraph("Power BI")

    document.add_paragraph("EDUCATION")
    document.add_paragraph(
        "Bachelor of Technology in Computer Science"
    )

    document.add_paragraph("WORK EXPERIENCE")
    document.add_paragraph("Data Analyst")
    document.add_paragraph("ABC Technologies")
    document.add_paragraph("2021 - Present")

    document.add_paragraph("PROJECTS")
    document.add_paragraph("Sales Dashboard")

    document.save(file_path)

    with open(file_path, "rb") as resume_file:
        response = client.post(
            "/api/analyze",
            data={
                "resume": (
                    resume_file,
                    "resume.docx",
                ),
                "job_description": """
                We are looking for a Data Analyst
                with 3+ years of experience.

                Bachelor's degree is required.

                Strong skills in Python, SQL, Pandas
                and Power BI.
                """,
            },
            content_type="multipart/form-data",
        )


    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"

    assert "ats_score" in data
    assert 0 <= data["ats_score"] <= 100

    assert "skills" in data
    assert "experience" in data
    assert "qualifications" in data
    assert "keywords" in data

    assert "component_scores" in data
    assert "weights" in data

    assert "matched_skills" in data
    assert "missing_skills" in data

    assert "matched_qualifications" in data
    assert "missing_qualifications" in data

    assert "matched_keywords" in data
    assert "missing_keywords" in data

    assert "score_explanation" in data

    assert "strengths" in data
    assert "weaknesses" in data

    assert "recommendations" in data

    assert "resume_quality" in data

    assert "improvement_suggestions" in data

    improvement_suggestions = data[
        "improvement_suggestions"
    ]

    assert isinstance(
        improvement_suggestions,
        list,
    )

    for suggestion in improvement_suggestions:
        assert "type" in suggestion
        assert "priority" in suggestion
        assert "current_text" in suggestion
        assert "suggestion" in suggestion
        assert "reason" in suggestion

    resume_quality = data["resume_quality"]

    assert "score" in resume_quality
    assert 0 <= resume_quality["score"] <= 100

    assert "components" in resume_quality
    assert "missing_sections" in resume_quality
    assert "optional_sections" in resume_quality

    assert "name" in resume_quality["components"]
    assert "email" in resume_quality["components"]
    assert "phone" in resume_quality["components"]
    assert "skills" in resume_quality["components"]
    assert "education" in resume_quality["components"]
    assert "experience" in resume_quality["components"]
    assert "projects" in resume_quality["components"]

    assert "certifications" in (
        resume_quality["optional_sections"]
    )

    assert "tailoring_plan" in response.get_json()
    assert isinstance(response.get_json()["tailoring_plan"], dict)


def test_ats_analysis_missing_resume():
    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    response = client.post(
        "/api/analyze",
        data={
            "job_description": "Python developer required."
        },
    )

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"


def test_ats_analysis_empty_job_description():
    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    resume_content = b"""
    Sai Chandra

    SKILLS
    Python
    SQL
    """

    response = client.post(
        "/api/analyze",
        data={
            "resume": (
                BytesIO(resume_content),
                "resume.txt",
            ),
            "job_description": "",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"


def test_ats_analysis_unsupported_resume_format():
    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    response = client.post(
        "/api/analyze",
        data={
            "resume": (
                BytesIO(b"fake resume"),
                "resume.txt",
            ),
            "job_description": "Python developer required.",
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"
    assert data["message"] == (
        "Unsupported file format. Use PDF or DOCX."
    )