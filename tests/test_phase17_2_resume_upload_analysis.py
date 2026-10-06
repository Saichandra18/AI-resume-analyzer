from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app import create_app


@pytest.fixture
def client():
    app = create_app()
    app.config["TESTING"] = True

    with app.test_client() as client:
        yield client


@pytest.fixture
def sample_docx_bytes():
    document = Document()

    document.add_heading("Alex Johnson", level=1)

    document.add_paragraph("alex@example.com")
    document.add_paragraph("+91-9876543210")
    document.add_paragraph("India")

    document.add_heading("Professional Summary", level=2)
    document.add_paragraph(
        "AI/ML Engineer with experience in Python, "
        "Machine Learning, SQL, Generative AI and RAG."
    )

    document.add_heading("Skills", level=2)
    document.add_paragraph(
        "Python, Machine Learning, SQL, Pandas, "
        "Generative AI, RAG, Flask"
    )

    document.add_heading("Experience", level=2)
    document.add_paragraph(
        "AI/ML Engineer - Example Technologies - 2024-2025"
    )

    document.add_heading("Projects", level=2)
    document.add_paragraph(
        "AI Resume Analyzer - Built a resume analysis "
        "and job-description matching system."
    )

    document.add_heading("Education", level=2)
    document.add_paragraph(
        "B.Tech - Example University - 2024"
    )

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output.getvalue()


@pytest.fixture
def sample_pdf_bytes():
    output = BytesIO()

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    writer.write(output)
    output.seek(0)

    return output.getvalue()


def test_upload_requires_resume_file(client):
    response = client.post("/api/resume/upload")

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"
    assert "resume" in data["message"].lower()


def test_upload_rejects_empty_filename(client):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(b""),
                "",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"
    assert "file" in data["message"].lower()


def test_upload_rejects_unsupported_file_type(client):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(b"plain text resume"),
                "resume.txt",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 400

    data = response.get_json()

    assert data["status"] == "error"
    assert "unsupported" in data["message"].lower()


def test_upload_accepts_docx(
    client,
    sample_docx_bytes,
):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_docx_bytes),
                "sample_resume.docx",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"
    assert data["message"] == "Resume uploaded successfully"
    assert data["filename"] == "sample_resume.docx"

    assert isinstance(data["text"], str)
    assert len(data["text"].strip()) > 0


def test_docx_upload_preserves_resume_content(
    client,
    sample_docx_bytes,
):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_docx_bytes),
                "content_resume.docx",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200

    text = response.get_json()["text"]

    assert "Alex Johnson" in text
    assert "alex@example.com" in text
    assert "Python" in text
    assert "Machine Learning" in text
    assert "Generative AI" in text
    assert "RAG" in text
    assert "B.Tech" in text


def test_upload_accepts_pdf(
    client,
    sample_pdf_bytes,
):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_pdf_bytes),
                "sample_resume.pdf",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"
    assert data["filename"] == "sample_resume.pdf"
    assert isinstance(data["text"], str)


def test_pdf_upload_returns_text_field(
    client,
    sample_pdf_bytes,
):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_pdf_bytes),
                "resume.pdf",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200

    data = response.get_json()

    assert "text" in data
    assert isinstance(data["text"], str)


def test_upload_response_has_expected_schema(
    client,
    sample_docx_bytes,
):
    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_docx_bytes),
                "schema_resume.docx",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200

    data = response.get_json()

    assert set(data.keys()) == {
        "status",
        "message",
        "filename",
        "text",
    }


def test_multiple_uploads_are_independent(
    client,
    sample_docx_bytes,
):
    first_response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_docx_bytes),
                "first_resume.docx",
            )
        },
        content_type="multipart/form-data",
    )

    second_response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                BytesIO(sample_docx_bytes),
                "second_resume.docx",
            )
        },
        content_type="multipart/form-data",
    )

    assert first_response.status_code == 200
    assert second_response.status_code == 200

    first_data = first_response.get_json()
    second_data = second_response.get_json()

    assert first_data["filename"] == "first_resume.docx"
    assert second_data["filename"] == "second_resume.docx"

    assert first_data["text"]
    assert second_data["text"]

def test_pdf_upload_extracts_resume_content(client):
    from reportlab.pdfgen import canvas

    output = BytesIO()

    pdf = canvas.Canvas(output)
    pdf.drawString(72, 750, "Alex Johnson")
    pdf.drawString(72, 730, "AI/ML Engineer")
    pdf.drawString(72, 710, "Python Machine Learning SQL")
    pdf.drawString(72, 690, "Generative AI RAG")
    pdf.save()

    output.seek(0)

    response = client.post(
        "/api/resume/upload",
        data={
            "resume": (
                output,
                "text_resume.pdf",
            )
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"

    text = data["text"]

    assert "Alex Johnson" in text
    assert "AI/ML Engineer" in text
    assert "Python" in text
    assert "Machine Learning" in text
    assert "Generative AI" in text
    assert "RAG" in text