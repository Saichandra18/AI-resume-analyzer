from pathlib import Path
from app.final_resume_models import build_final_resume

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


from app.final_resume_models import FinalResume
from app.resume_document_renderer import (
    render_final_resume,
)


def make_resume():
    return FinalResume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
            "phone": "9876543210",
            "location": "India",
        },
        summary=(
            "Machine learning engineer with "
            "Python experience."
        ),
        skills=[
            "Python",
            "Machine Learning",
            "SQL",
        ],
        experience=[
            {
                "job_title": "ML Engineer",
                "company": "Example Corp",
                "dates": "2024 - Present",
                "bullets": [
                    "Built machine learning models.",
                    "Improved data processing workflows.",
                ],
            }
        ],
        projects=[
            {
                "name": "Resume Analyzer",
                "description": (
                    "Built an AI-powered resume analyzer."
                ),
            }
        ],
        education=[
            {
                "degree": "B.Tech",
                "institution": "Example University",
                "year": "2024",
            }
        ],
        certifications=[
            {
                "name": "AWS Certified Cloud Practitioner"
            }
        ],
    )


def test_render_final_resume_creates_docx(tmp_path):
    output = (
        tmp_path / "resume.docx"
    )

    result = render_final_resume(
        make_resume(),
        output,
    )

    assert result == output
    assert output.exists()
    assert output.stat().st_size > 0


def test_generated_docx_can_be_opened(tmp_path):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    assert len(document.paragraphs) > 0


def test_generated_docx_contains_candidate_name(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "John Doe" in text


def test_generated_docx_contains_contact_information(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "john@example.com" in text
    assert "9876543210" in text


def test_generated_docx_contains_sections(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Professional Summary" in text
    assert "Skills" in text
    assert "Experience" in text
    assert "Projects" in text
    assert "Education" in text
    assert "Certifications" in text


def test_generated_docx_contains_experience_content(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "ML Engineer" in text
    assert "Example Corp" in text
    assert "Built machine learning models." in text


def test_generated_docx_contains_project_content(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Resume Analyzer" in text
    assert (
        "Built an AI-powered resume analyzer."
        in text
    )


def test_generated_docx_contains_skills(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Python" in text
    assert "Machine Learning" in text
    assert "SQL" in text


def test_generated_docx_respects_section_order(
    tmp_path,
):
    resume = make_resume()

    resume.section_order = [
        "projects",
        "skills",
        "experience",
        "education",
        "certifications",
        "summary",
    ]

    output = (
        tmp_path / "ordered.docx"
    )

    render_final_resume(
        resume,
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    positions = [
        text.index("Projects"),
        text.index("Skills"),
        text.index("Experience"),
        text.index("Education"),
        text.index("Certifications"),
        text.index("Professional Summary"),
    ]

    assert positions == sorted(
        positions
    )


def test_empty_optional_sections_are_not_rendered(
    tmp_path,
):
    resume = FinalResume(
        personal_info={
            "name": "John Doe"
        },
        summary="Engineer.",
    )

    output = (
        tmp_path / "minimal.docx"
    )

    render_final_resume(
        resume,
        output,
    )

    document = Document(
        str(output)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Professional Summary" in text
    assert "Experience" not in text
    assert "Projects" not in text
    assert "Certifications" not in text


def test_invalid_resume_is_rejected(tmp_path):
    resume = FinalResume()

    output = (
        tmp_path / "invalid.docx"
    )

    with pytest.raises(ValueError):
        render_final_resume(
            resume,
            output,
        )


def test_non_final_resume_is_rejected(tmp_path):
    output = (
        tmp_path / "invalid.docx"
    )

    with pytest.raises(TypeError):
        render_final_resume(
            {},
            output,
        )


def test_non_docx_output_is_rejected(tmp_path):
    output = (
        tmp_path / "resume.pdf"
    )

    with pytest.raises(ValueError):
        render_final_resume(
            make_resume(),
            output,
        )


def test_nested_output_directory_is_created(
    tmp_path,
):
    output = (
        tmp_path
        / "generated"
        / "resumes"
        / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    assert output.exists()


def test_renderer_does_not_mutate_resume(tmp_path):
    resume = make_resume()

    before = resume.to_dict()

    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        resume,
        output,
    )

    assert resume.to_dict() == before


def test_validation_can_be_disabled(tmp_path):
    resume = FinalResume()

    output = (
        tmp_path / "unvalidated.docx"
    )

    render_final_resume(
        resume,
        output,
        validate=False,
    )

    assert output.exists()


def test_document_has_reasonable_margins(
    tmp_path,
):
    output = (
        tmp_path / "resume.docx"
    )

    render_final_resume(
        make_resume(),
        output,
    )

    document = Document(
        str(output)
    )

    section = document.sections[0]

    assert (
        section.top_margin.inches
        == pytest.approx(0.55)
    )

    assert (
        section.bottom_margin.inches
        == pytest.approx(0.55)
    )

def test_format_paragraph_applies_layout_settings():
    from docx import Document
    from app.resume_document_renderer import _format_paragraph

    document = Document()
    paragraph = document.add_paragraph()

    _format_paragraph(
        paragraph,
        space_before=8,
        space_after=3,
        line_spacing=1.0,
        keep_with_next=True,
        keep_together=True,
    )

    fmt = paragraph.paragraph_format

    assert fmt.space_before.pt == 8
    assert fmt.space_after.pt == 3
    assert fmt.line_spacing == 1.0
    assert fmt.keep_with_next is True
    assert fmt.keep_together is True


def test_section_heading_keeps_with_next(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="Experienced engineer.",
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Professional Summary"
    )

    assert heading.paragraph_format.keep_with_next is True
    assert heading.paragraph_format.space_before.pt == 8
    assert heading.paragraph_format.space_after.pt == 3


def test_name_uses_expected_typography(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    name_paragraph = document.paragraphs[0]
    run = name_paragraph.runs[0]

    assert run.text == "John Doe"
    assert run.bold is True
    assert run.font.name == "Arial"
    assert run.font.size.pt == 18
    assert name_paragraph.paragraph_format.keep_with_next is True


def test_contact_line_uses_compact_spacing(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
            "phone": "1234567890",
        },
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    contact = document.paragraphs[1]

    assert "john@example.com" in contact.text
    assert "1234567890" in contact.text
    assert contact.paragraph_format.keep_with_next is True
    assert contact.paragraph_format.space_after.pt == 6


def test_item_title_keeps_with_secondary_information(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "company": "ABC Technologies",
                "dates": "2024-2026",
                "bullets": ["Built ML pipelines."],
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    title = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "ML Engineer"
    )

    assert title.paragraph_format.keep_with_next is True
    assert title.paragraph_format.keep_together is True


def test_bullet_paragraph_keeps_together(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "bullets": ["Built ML pipelines."],
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    bullet = next(
        paragraph
        for paragraph in document.paragraphs
        if "Built ML pipelines." in paragraph.text
    )

    assert bullet.paragraph_format.keep_together is True
    assert bullet.paragraph_format.space_after.pt == 3


def test_bullets_take_priority_over_description(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "description": "Built an ML pipeline.",
                "bullets": ["Built an ML pipeline."],
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert full_text.count("Built an ML pipeline.") == 1


def test_description_is_used_when_bullets_are_absent(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "description": "Built an ML pipeline.",
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert full_text.count("Built an ML pipeline.") == 1


def test_multiple_bullets_remain_separate(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "bullets": [
                    "Built ML pipelines.",
                    "Improved model performance.",
                    "Deployed models to production.",
                ],
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    bullet_texts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Bullet"
    ]

    assert bullet_texts == [
        "Built ML pipelines.",
        "Improved model performance.",
        "Deployed models to production.",
    ]

def test_skills_list_content_is_preserved(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        skills=[
            "Python",
            "Java",
            "Scikit-learn",
            "TensorFlow",
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Python" in full_text
    assert "Java" in full_text
    assert "Scikit-learn" in full_text
    assert "TensorFlow" in full_text

def test_education_content_is_preserved(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        education=[
            {
                "degree": "B.Tech",
                "institution": "ABC University",
                "year": "2024",
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "B.Tech" in full_text
    assert "ABC University" in full_text
    assert "2024" in full_text


def test_certification_content_is_preserved(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        certifications=[
            {
                "name": "AWS Certified Cloud Practitioner",
                "organization": "Amazon Web Services",
                "year": "2025",
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "AWS Certified Cloud Practitioner" in full_text
    assert "Amazon Web Services" in full_text
    assert "2025" in full_text


def test_unicode_content_is_preserved(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={
            "name": "José García",
            "location": "São Paulo",
        },
        summary="NLP – Deep Learning with C++ and C#.",
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "José García" in full_text
    assert "São Paulo" in full_text
    assert "NLP – Deep Learning with C++ and C#." in full_text


def test_item_without_optional_secondary_fields_renders(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "bullets": [
                    "Built machine learning models."
                ],
            }
        ],
    )

    output_path = tmp_path / "resume.docx"
    render_final_resume(resume, output_path)

    document = Document(output_path)

    full_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "ML Engineer" in full_text
    assert "Built machine learning models." in full_text

def test_existing_bullet_markers_are_not_duplicated(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "bullets": [
                    "• Built ML pipelines.",
                    " Improved model performance.",
                    "- Deployed models to production.",
                ],
            }
        ],
    )

    output_path = tmp_path / "bullet_cleanup.docx"

    render_final_resume(
        resume,
        output_path,
    )

    document = Document(output_path)

    bullet_texts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Bullet"
    ]

    assert bullet_texts == [
        "Built ML pipelines.",
        "Improved model performance.",
        "Deployed models to production.",
    ]

def test_renderer_accepts_ats_clean_template(tmp_path):
    """Renderer should accept the ATS Clean template."""
    output = tmp_path / "ats_clean.docx"

    result = render_final_resume(
        make_resume(),
        output,
        template="ats_clean",
    )

    assert result == output
    assert output.exists()


def test_renderer_defaults_to_default_template(tmp_path):
    """Renderer should preserve the existing default behavior."""
    output = tmp_path / "default.docx"

    result = render_final_resume(
        make_resume(),
        output,
    )

    assert result == output
    assert output.exists()


def test_renderer_rejects_unknown_template(tmp_path):
    """Renderer should reject unknown template names."""
    output = tmp_path / "invalid.docx"

    with pytest.raises(ValueError, match="Unknown resume template"):
        render_final_resume(
            make_resume(),
            output,
            template="does-not-exist",
        )

def test_renderer_applies_template_document_defaults(
    tmp_path,
):
    """Renderer should apply document settings from the template."""

    output_path = tmp_path / "ats_clean_defaults.docx"

    render_final_resume(
        make_resume(),
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)
    section = document.sections[0]

    assert section.top_margin.inches == pytest.approx(0.55)
    assert section.bottom_margin.inches == pytest.approx(0.55)
    assert section.left_margin.inches == pytest.approx(0.65)
    assert section.right_margin.inches == pytest.approx(0.65)

    normal = document.styles["Normal"]

    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 10

def test_name_uses_template_configuration(tmp_path):
    """Name formatting should come from the selected template."""

    output_path = tmp_path / "template_name.docx"

    render_final_resume(
        make_resume(),
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    name_paragraph = document.paragraphs[0]
    run = name_paragraph.runs[0]

    assert run.text == "John Doe"
    assert run.bold is True
    assert run.font.name == "Arial"
    assert run.font.size.pt == 18
    assert (
        name_paragraph.alignment
        == WD_ALIGN_PARAGRAPH.CENTER
    )


def test_contact_uses_template_configuration(tmp_path):
    """Contact formatting should come from the selected template."""

    output_path = tmp_path / "template_contact.docx"

    render_final_resume(
        make_resume(),
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    contact = document.paragraphs[1]

    assert "john@example.com" in contact.text
    assert "9876543210" in contact.text
    assert (
        contact.alignment
        == WD_ALIGN_PARAGRAPH.CENTER
    )
    assert contact.paragraph_format.keep_with_next is True
    assert contact.paragraph_format.space_after.pt == 6

    run = contact.runs[0]

    assert run.font.name == "Arial"
    assert run.font.size.pt == 9

def test_section_heading_uses_template_configuration(
    tmp_path,
):
    """Section headings should use template configuration."""

    output_path = tmp_path / "template_heading.docx"

    render_final_resume(
        make_resume(),
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Professional Summary"
    )

    run = heading.runs[0]

    assert run.text == "Professional Summary"
    assert run.bold is True
    assert run.font.name == "Arial"
    assert run.font.size.pt == 12
    assert heading.paragraph_format.space_before.pt == 8
    assert heading.paragraph_format.space_after.pt == 3
    assert heading.paragraph_format.line_spacing == 1.0
    assert heading.paragraph_format.keep_with_next is True

def test_ats_clean_template_uses_ats_safe_configuration():
    from app.templates.registry import get_template

    template = get_template("ats_clean")

    assert template.ats_safe is True
    assert template.single_column is True
    assert template.use_bullets is True
    assert template.font_name == "Arial"


def test_ats_clean_template_uses_compact_typography():
    from app.templates.registry import get_template

    template = get_template("ats_clean")

    assert template.body_font_size == 10
    assert template.heading_font_size == 12
    assert template.name_font_size == 18
    assert template.contact_font_size == 9
    assert template.item_title_font_size == 10.5
    assert template.secondary_font_size == 9


def test_ats_clean_template_uses_compact_margins():
    from app.templates.registry import get_template

    template = get_template("ats_clean")

    assert template.top_margin == 0.55
    assert template.bottom_margin == 0.55
    assert template.left_margin == 0.65
    assert template.right_margin == 0.65


def test_ats_clean_template_renders_complete_resume(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
            "phone": "9876543210",
        },
        summary="AI/ML Engineer specializing in Generative AI.",
        skills=[
            "Python",
            "SQL",
            "Machine Learning",
            "RAG",
        ],
        experience=[
            {
                "job_title": "AI Engineer",
                "company": "Example Company",
                "dates": "2025 - Present",
                "bullets": [
                    "Built machine learning pipelines.",
                    "Developed RAG applications.",
                ],
            }
        ],
        projects=[
            {
                "title": "AI Resume Analyzer",
                "description": [
                    "Built an ATS analysis platform."
                ],
            }
        ],
    )

    output_path = tmp_path / "ats_clean.docx"

    render_final_resume(
        resume,
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "John Doe" in text
    assert "Professional Summary" in text
    assert "Skills" in text
    assert "Experience" in text
    assert "Projects" in text
    assert "Python" in text
    assert "Built machine learning pipelines." in text


def test_ats_clean_template_does_not_create_tables(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
        skills=["Python", "SQL"],
        experience=[
            {
                "job_title": "AI Engineer",
                "bullets": ["Built ML systems."],
            }
        ],
    )

    output_path = tmp_path / "ats_clean.docx"

    render_final_resume(
        resume,
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    assert len(document.tables) == 0


def test_ats_clean_template_uses_single_column_document(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
    )

    output_path = tmp_path / "ats_clean.docx"

    render_final_resume(
        resume,
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    # The renderer currently uses normal paragraphs only.
    # Therefore no table-based multi-column layout is introduced.
    assert len(document.tables) == 0


def test_ats_clean_template_bullets_remain_separate(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "AI Engineer",
                "bullets": [
                    "Built ML pipelines.",
                    "Developed RAG systems.",
                    "Improved ATS matching.",
                ],
            }
        ],
    )

    output_path = tmp_path / "ats_clean.docx"

    render_final_resume(
        resume,
        output_path,
        template="ats_clean",
    )

    document = Document(output_path)

    bullet_texts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text in {
            "Built ML pipelines.",
            "Developed RAG systems.",
            "Improved ATS matching.",
        }
    ]

    assert bullet_texts == [
        "Built ML pipelines.",
        "Developed RAG systems.",
        "Improved ATS matching.",
    ]

def test_renderer_accepts_modern_template(tmp_path):
    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    assert output_path.exists()


def test_modern_template_name_is_left_aligned(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    name_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "John Doe"
    )

    assert (
        name_paragraph.alignment
        == WD_ALIGN_PARAGRAPH.LEFT
    )


def test_modern_template_name_uses_expected_typography(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    name_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "John Doe"
    )

    run = name_paragraph.runs[0]

    assert run.font.name == "Aptos"
    assert run.font.size.pt == 20
    assert run.bold is True


def test_modern_template_contact_is_left_aligned(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    resume = build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
        },
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    contact_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "john@example.com" in paragraph.text
    )

    assert (
        contact_paragraph.alignment
        == WD_ALIGN_PARAGRAPH.LEFT
    )


def test_modern_template_section_heading_uses_expected_typography(
    tmp_path,
):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Professional Summary"
    )

    run = heading.runs[0]

    assert run.font.name == "Aptos"
    assert run.font.size.pt == 13
    assert run.bold is True


def test_modern_template_item_title_uses_expected_typography(
    tmp_path,
):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "AI Engineer",
                "company": "Example Company",
                "bullets": [
                    "Built AI systems.",
                ],
            }
        ],
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    title = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "AI Engineer"
    )

    run = title.runs[0]

    assert run.font.name == "Aptos"
    assert run.font.size.pt == 11
    assert run.bold is True


def test_modern_template_renders_complete_resume(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
            "phone": "9876543210",
        },
        summary="AI/ML Engineer specializing in Generative AI.",
        skills=[
            "Python",
            "SQL",
            "Machine Learning",
            "RAG",
        ],
        experience=[
            {
                "job_title": "AI Engineer",
                "company": "Example Company",
                "dates": "2025 - Present",
                "bullets": [
                    "Built machine learning pipelines.",
                    "Developed RAG applications.",
                ],
            }
        ],
        projects=[
            {
                "title": "AI Resume Analyzer",
                "description": [
                    "Built an ATS analysis platform.",
                ],
            }
        ],
        education=[
            {
                "degree": "B.Tech",
                "institution": "Example University",
                "year": "2025",
            }
        ],
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "John Doe" in text
    assert "Professional Summary" in text
    assert "Skills" in text
    assert "Experience" in text
    assert "Projects" in text
    assert "Education" in text
    assert "Python" in text
    assert "Built machine learning pipelines." in text


def test_modern_template_does_not_create_tables(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
        skills=["Python", "SQL"],
        experience=[
            {
                "job_title": "AI Engineer",
                "bullets": [
                    "Built AI systems.",
                ],
            }
        ],
    )

    output_path = tmp_path / "modern.docx"

    render_final_resume(
        resume,
        output_path,
        template="modern",
    )

    document = Document(output_path)

    assert len(document.tables) == 0
    
def test_renderer_accepts_technical_template(tmp_path):
    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    assert output_path.exists()


def test_technical_template_name_is_left_aligned(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    name_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "John Doe"
    )

    assert (
        name_paragraph.alignment
        == WD_ALIGN_PARAGRAPH.LEFT
    )


def test_technical_template_name_uses_expected_typography(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    name_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "John Doe"
    )

    run = name_paragraph.runs[0]

    assert run.font.name == "Arial"
    assert run.font.size.pt == 17
    assert run.bold is True


def test_technical_template_contact_is_left_aligned(tmp_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    resume = build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
        },
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    contact_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if "john@example.com" in paragraph.text
    )

    assert (
        contact_paragraph.alignment
        == WD_ALIGN_PARAGRAPH.LEFT
    )


def test_technical_template_section_heading_uses_expected_typography(
    tmp_path,
):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Professional Summary"
    )

    run = heading.runs[0]

    assert run.font.name == "Arial"
    assert run.font.size.pt == 11.5
    assert run.bold is True


def test_technical_template_item_title_uses_expected_typography(
    tmp_path,
):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        experience=[
            {
                "job_title": "ML Engineer",
                "company": "Example Company",
                "bullets": [
                    "Built ML pipelines.",
                ],
            }
        ],
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    title = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "ML Engineer"
    )

    run = title.runs[0]

    assert run.font.name == "Arial"
    assert run.font.size.pt == 10
    assert run.bold is True


def test_technical_template_renders_complete_resume(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
            "phone": "9876543210",
        },
        summary="AI/ML Engineer specializing in Generative AI.",
        skills=[
            "Python",
            "SQL",
            "Machine Learning",
            "RAG",
            "PySpark",
        ],
        experience=[
            {
                "job_title": "ML Engineer",
                "company": "Example Company",
                "dates": "2025 - Present",
                "bullets": [
                    "Built machine learning pipelines.",
                    "Developed RAG applications.",
                ],
            }
        ],
        projects=[
            {
                "title": "AI Resume Analyzer",
                "description": [
                    "Built an ATS analysis platform.",
                ],
            }
        ],
        education=[
            {
                "degree": "B.Tech",
                "institution": "Example University",
                "year": "2025",
            }
        ],
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "John Doe" in text
    assert "Professional Summary" in text
    assert "Skills" in text
    assert "Experience" in text
    assert "Projects" in text
    assert "Education" in text
    assert "Python" in text
    assert "PySpark" in text
    assert "Built machine learning pipelines." in text


def test_technical_template_does_not_create_tables(tmp_path):
    from docx import Document

    resume = build_final_resume(
        personal_info={"name": "John Doe"},
        summary="AI Engineer",
        skills=[
            "Python",
            "SQL",
            "PySpark",
        ],
        experience=[
            {
                "job_title": "ML Engineer",
                "bullets": [
                    "Built ML systems.",
                ],
            }
        ],
    )

    output_path = tmp_path / "technical.docx"

    render_final_resume(
        resume,
        output_path,
        template="technical",
    )

    document = Document(output_path)

    assert len(document.tables) == 0