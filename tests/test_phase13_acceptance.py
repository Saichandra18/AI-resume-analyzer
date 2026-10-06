from pathlib import Path

from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.rendered_document_validator import (
    validate_rendered_documents,
)


def _create_acceptance_docx(path):
    """Create a representative optimized resume DOCX."""
    document = Document()

    document.add_heading(
        "Sai Chandra Mudda",
        level=1,
    )

    document.add_paragraph(
        "AI/ML Engineer specializing in Python, "
        "Machine Learning, Generative AI and RAG."
    )

    document.add_heading(
        "Skills",
        level=2,
    )

    document.add_paragraph(
        "Python | SQL | Machine Learning | "
        "Generative AI | LLM | RAG"
    )

    document.add_heading(
        "Projects",
        level=2,
    )

    document.add_paragraph(
        "AI Resume Analyzer"
    )

    document.save(path)


def _create_acceptance_pdf(path):
    """Create a representative one-page optimized resume PDF."""
    pdf = canvas.Canvas(str(path))

    pdf.drawString(
        72,
        750,
        "Sai Chandra Mudda",
    )

    pdf.drawString(
        72,
        720,
        "AI/ML Engineer specializing in Python, "
        "Machine Learning, Generative AI and RAG.",
    )

    pdf.drawString(
        72,
        690,
        "Skills: Python | SQL | Machine Learning | "
        "Generative AI | LLM | RAG",
    )

    pdf.drawString(
        72,
        660,
        "Projects: AI Resume Analyzer",
    )

    pdf.save()


def test_phase13_final_acceptance(
    tmp_path,
):
    """
    Final Phase 13 acceptance test.

    Verifies that optimized resume artifacts:
    - exist,
    - are readable,
    - contain resume content,
    - satisfy the requested PDF page count,
    - pass structural validation.
    """

    docx_path = (
        tmp_path / "optimized_resume.docx"
    )

    pdf_path = (
        tmp_path / "optimized_resume.pdf"
    )

    # ---------------------------------------------------------
    # Generate representative final artifacts.
    # ---------------------------------------------------------

    _create_acceptance_docx(docx_path)
    _create_acceptance_pdf(pdf_path)

    # ---------------------------------------------------------
    # Verify files exist.
    # ---------------------------------------------------------

    assert docx_path.exists()
    assert pdf_path.exists()

    assert docx_path.stat().st_size > 0
    assert pdf_path.stat().st_size > 0

    # ---------------------------------------------------------
    # Validate rendered documents.
    # ---------------------------------------------------------

    validation = validate_rendered_documents(
        docx_path=docx_path,
        pdf_path=pdf_path,
        target_pages=1,
    )

    assert validation["valid"] is True

    # ---------------------------------------------------------
    # PDF acceptance checks.
    # ---------------------------------------------------------

    assert validation["pdf"]["valid"] is True
    assert validation["pdf"]["readable"] is True
    assert validation["pdf"]["text_present"] is True

    assert (
        validation["pdf"]["page_count"]["valid"]
        is True
    )

    assert (
        validation["pdf"]["page_count"]["target_pages"]
        == 1
    )

    assert (
        validation["pdf"]["page_count"]["actual_pages"]
        == 1
    )

    # ---------------------------------------------------------
    # DOCX acceptance checks.
    # ---------------------------------------------------------

    assert validation["docx"]["valid"] is True
    assert validation["docx"]["readable"] is True
    assert validation["docx"]["text_present"] is True

    # ---------------------------------------------------------
    # Independently verify DOCX content.
    # ---------------------------------------------------------

    document = Document(str(docx_path))

    docx_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Sai Chandra Mudda" in docx_text
    assert "AI/ML Engineer" in docx_text
    assert "Python" in docx_text
    assert "Machine Learning" in docx_text
    assert "Generative AI" in docx_text
    assert "RAG" in docx_text
    assert "AI Resume Analyzer" in docx_text

    # ---------------------------------------------------------
    # Independently verify PDF content.
    # ---------------------------------------------------------

    reader = PdfReader(str(pdf_path))

    assert len(reader.pages) == 1

    pdf_text = "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )

    assert "Sai Chandra Mudda" in pdf_text
    assert "AI/ML Engineer" in pdf_text
    assert "Python" in pdf_text
    assert "Machine Learning" in pdf_text
    assert "Generative AI" in pdf_text
    assert "RAG" in pdf_text
    assert "AI Resume Analyzer" in pdf_text