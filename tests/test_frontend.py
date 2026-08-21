from app import create_app


def test_home_page_returns_html():
    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    response = client.get("/")

    assert response.status_code == 200
    assert response.content_type.startswith("text/html")
    assert b"AI Resume Analyzer" in response.data

def test_analyze_api_returns_frontend_data(tmp_path):
    from docx import Document

    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    file_path = tmp_path / "resume.docx"

    document = Document()

    document.add_paragraph("Sai Chandra")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python")
    document.add_paragraph("SQL")
    document.add_paragraph("EDUCATION")
    document.add_paragraph(
        "Bachelor of Technology"
    )

    document.save(file_path)

    with open(file_path, "rb") as resume_file:
        response = client.post(
            "/api/analyze",
            data={
                "resume": (
                    resume_file,
                    "resume.docx",
                ),
                "job_description": """
                Python developer with 2+ years
                experience. Bachelor's degree required.
                Python and SQL skills required.
                """,
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"
    assert "ats_score" in data
    assert "matched_skills" in data
    assert "missing_skills" in data
    assert "score_explanation" in data
    assert "strengths" in data
    assert "weaknesses" in data
    assert "recommendations" in data