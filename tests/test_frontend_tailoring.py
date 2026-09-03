from pathlib import Path

from app import create_app


BASE_DIR = Path(__file__).resolve().parents[1]

INDEX_PATH = (
    BASE_DIR
    / "frontend"
    / "templates"
    / "index.html"
)

APP_JS_PATH = (
    BASE_DIR
    / "frontend"
    / "static"
    / "js"
    / "app.js"
)


def test_home_page_contains_tailoring_plan_section():
    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    response = client.get("/")

    assert response.status_code == 200

    assert b"resume-tailoring" in response.data
    assert b"Resume Tailoring Plan" in response.data


def test_home_page_contains_tailoring_plan_fields():
    html = INDEX_PATH.read_text(
        encoding="utf-8"
    )

    assert 'id="resume-tailoring"' in html
    assert 'class="resume-tailoring"' in html
    assert 'id="tailoring-priorities"' in html
    assert 'id="tailoring-summary"' in html
    assert 'id="tailoring-skills"' in html
    assert 'id="tailoring-experience"' in html
    assert 'id="tailoring-keywords"' in html
    assert 'id="tailoring-qualifications"' in html
    assert 'id="tailoring-warnings"' in html


def test_app_js_contains_tailoring_plan_renderer():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert "displayTailoringPlan" in javascript
    assert "tailoring_plan" in javascript


def test_app_js_uses_tailoring_plan_data():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert "data.tailoring_plan" in javascript
    assert "plan.priorities" in javascript
    assert "plan.summary" in javascript
    assert "plan.skills" in javascript
    assert "plan.experience" in javascript
    assert "plan.keywords" in javascript
    assert "plan.qualifications" in javascript
    assert "plan.warnings" in javascript


def test_home_page_contains_tailoring_plan_labels():
    html = INDEX_PATH.read_text(
        encoding="utf-8"
    )

    assert "Priorities" in html
    assert "Summary" in html
    assert "Skills" in html
    assert "Experience" in html
    assert "Keywords" in html
    assert "Qualifications" in html
    assert "Warnings" in html

def test_app_js_uses_tailoring_item_fields():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert "item.priority" in javascript
    assert "item.action" in javascript
    assert "item.recommendation" in javascript


def test_app_js_hides_empty_tailoring_sections():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert (
        'section.items.length > 0'
        in javascript
    )

    assert (
        'sectionWrapper.style.display'
        in javascript
    )

def test_app_js_handles_tailoring_warnings():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert "plan.warnings" in javascript
    assert "warningsContainer" in javascript
    assert "warning" in javascript

def test_app_js_hides_empty_tailoring_sections():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert (
        'section.items.length > 0'
        in javascript
    )

    assert (
        'sectionWrapper.style.display'
        in javascript
    )


def test_app_js_calls_tailoring_plan_renderer():
    javascript = APP_JS_PATH.read_text(
        encoding="utf-8"
    )

    assert "displayTailoringPlan(data);" in javascript


def test_analyze_api_returns_tailoring_plan(tmp_path):
    from docx import Document

    app = create_app()
    app.config["TESTING"] = True

    client = app.test_client()

    file_path = tmp_path / "test_tailoring_resume.docx"

    document = Document()

    document.add_paragraph("Sai Chandra")
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python")
    document.add_paragraph("SQL")
    document.add_paragraph("EDUCATION")
    document.add_paragraph(
        "Bachelor of Technology"
    )

    document.save(file_path)

    with open(file_path, "rb") as resume_file:
        response = client.post(
            "/api/analyze",
            data={
                "resume": (
                    resume_file,
                    "resume.docx",
                ),
                "job_description": """
                Python developer with 2+ years
                experience. Bachelor's degree required.
                Python and SQL skills required.
                Machine Learning experience preferred.
                """,
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"
    assert "tailoring_plan" in data
    assert isinstance(
        data["tailoring_plan"],
        dict,
    )

    assert "priorities" in data["tailoring_plan"]
    assert "summary" in data["tailoring_plan"]
    assert "skills" in data["tailoring_plan"]
    assert "experience" in data["tailoring_plan"]
    assert "keywords" in data["tailoring_plan"]
    assert "qualifications" in data["tailoring_plan"]
    assert "warnings" in data["tailoring_plan"]