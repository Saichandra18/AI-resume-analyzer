from app.services.resume_parser import extract_structured_resume


def test_extract_structured_resume_from_docx(tmp_path):
    from docx import Document

    file_path = tmp_path / "resume.docx"

    document = Document()

    document.add_paragraph("Sai Chandra")
    document.add_paragraph("sai@example.com")
    document.add_paragraph("+91 9876543210")

    document.add_paragraph("SKILLS")
    document.add_paragraph("Python")
    document.add_paragraph("SQL")
    document.add_paragraph("Power BI")

    document.add_paragraph("EDUCATION")
    document.add_paragraph("Bachelor of Technology in Computer Science")
    document.add_paragraph("ABC University")
    document.add_paragraph("2020 - 2024")

    document.add_paragraph("WORK EXPERIENCE")
    document.add_paragraph("Data Analyst")
    document.add_paragraph("ABC Technologies")
    document.add_paragraph("2024 - Present")

    document.add_paragraph("PROJECTS")
    document.add_paragraph("AI Resume Analyzer")

    document.add_paragraph("CERTIFICATIONS")
    document.add_paragraph("AWS Certified Cloud Practitioner")

    document.save(file_path)

    resume = extract_structured_resume(file_path)

    assert resume["personal_info"]["name"] == "Sai Chandra"
    assert resume["personal_info"]["email"] == "sai@example.com"
    assert resume["personal_info"]["phone"] == "+91 9876543210"

    assert "Python" in resume["skills"]
    assert "SQL" in resume["skills"]
    assert "Power BI" in resume["skills"]

    assert "Bachelor of Technology in Computer Science" in resume["education"]["degrees"]

    assert "Data Analyst" in resume["experience"]["job_titles"]

    assert "AI Resume Analyzer" in resume["projects"]

    assert "AWS Certified Cloud Practitioner" in resume["certifications"]