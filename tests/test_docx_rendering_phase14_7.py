from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.final_resume_models import build_final_resume
from app.resume_document_renderer import render_final_resume
from app.templates.registry import get_template


TEMPLATES = ["ats_clean", "modern", "technical"]


def make_phase14_7_resume():
    return build_final_resume(
        personal_info={
            "name": "John Doe",
            "email": "john@example.com",
            "phone": "9876543210",
            "location": "India",
        },
        summary="Machine learning engineer with Python experience.",
        skills=["Python", "Machine Learning", "SQL"],
        experience=[
            {
                "job_title": "ML Engineer",
                "company": "Example Corp",
                "dates": "2024 - Present",
                "bullets": [
                    "Built machine learning models.",
                    "Improved data processing workflows.",
                ],
            }
        ],
        projects=[
            {
                "name": "Resume Analyzer",
                "description": "Built an AI-powered resume analyzer.",
            }
        ],
        education=[
            {
                "degree": "B.Tech",
                "institution": "Example University",
                "year": "2024",
            }
        ],
        certifications=[
            {
                "name": "AWS Certified Cloud Practitioner",
            }
        ],
    )


def document_text(document):
    return "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )


def test_all_templates_generate_openable_docx(tmp_path):
    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        result = render_final_resume(
            make_phase14_7_resume(),
            output,
            template=template,
        )

        assert result == output
        assert output.exists()
        assert output.stat().st_size > 0

        document = Document(output)

        assert len(document.paragraphs) > 0


def test_all_templates_preserve_resume_content(tmp_path):
    resume = make_phase14_7_resume()

    expected = [
        "John Doe",
        "john@example.com",
        "9876543210",
        "India",
        "Machine learning engineer with Python experience.",
        "Python",
        "Machine Learning",
        "SQL",
        "ML Engineer",
        "Example Corp",
        "Built machine learning models.",
        "Improved data processing workflows.",
        "Resume Analyzer",
        "Built an AI-powered resume analyzer.",
        "B.Tech",
        "Example University",
        "2024",
        "AWS Certified Cloud Practitioner",
    ]

    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            resume,
            output,
            template=template,
        )

        document = Document(output)
        text = document_text(document)

        for value in expected:
            assert value in text


def test_all_templates_preserve_section_order(tmp_path):
    resume = make_phase14_7_resume()

    expected_order = [
        "Professional Summary",
        "Skills",
        "Experience",
        "Projects",
        "Education",
        "Certifications",
    ]

    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            resume,
            output,
            template=template,
        )

        text = document_text(Document(output))

        positions = [
            text.index(section)
            for section in expected_order
        ]

        assert positions == sorted(positions)


def test_all_templates_are_single_column_without_tables(tmp_path):
    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            make_phase14_7_resume(),
            output,
            template=template,
        )

        document = Document(output)

        assert len(document.tables) == 0

        for paragraph in document.paragraphs:
            assert paragraph.text is not None


def test_template_alignment_matches_configuration(tmp_path):
    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            make_phase14_7_resume(),
            output,
            template=template,
        )

        document = Document(output)
        config = get_template(template)

        name_paragraph = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text == "John Doe"
        )

        expected_alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[config.name_alignment]

        assert name_paragraph.alignment == expected_alignment


def test_template_typography_is_applied_to_name(tmp_path):
    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            make_phase14_7_resume(),
            output,
            template=template,
        )

        document = Document(output)
        config = get_template(template)

        name_paragraph = next(
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text == "John Doe"
        )

        run = name_paragraph.runs[0]

        assert run.font.size.pt == config.name_font_size
        assert run.bold is True


def test_rendering_does_not_mutate_final_resume(tmp_path):
    resume = make_phase14_7_resume()
    before = deepcopy(resume)

    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            resume,
            output,
            template=template,
        )

    assert resume.personal_info == before.personal_info
    assert resume.summary == before.summary
    assert resume.skills == before.skills
    assert resume.experience == before.experience
    assert resume.projects == before.projects
    assert resume.education == before.education
    assert resume.certifications == before.certifications
    assert resume.section_order == before.section_order

    assert resume.metadata.__dict__ == before.metadata.__dict__


def test_bullets_remain_separate_paragraphs_for_all_templates(tmp_path):
    for template in TEMPLATES:
        output = tmp_path / f"{template}.docx"

        render_final_resume(
            make_phase14_7_resume(),
            output,
            template=template,
        )

        document = Document(output)

        bullet_paragraphs = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text in {
                "Built machine learning models.",
                "Improved data processing workflows.",
            }
        ]

        assert len(bullet_paragraphs) == 2
        assert bullet_paragraphs[0].text != bullet_paragraphs[1].text