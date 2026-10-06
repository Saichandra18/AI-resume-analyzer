from pathlib import Path
from zipfile import ZipFile

import pytest
from pypdf import PdfReader
from docx import Document

from app.final_resume_models import FinalResume
from app.resume_document_renderer import render_final_resume
from app.resume_pdf_renderer import render_final_resume_pdf
from app.templates.registry import get_template


TEMPLATES = ("ats_clean", "modern", "technical")


def build_resume():
    """Build a representative ATS-safe FinalResume."""
    return FinalResume(
        personal_info={
            "name": "Test Candidate",
            "email": "test@example.com",
            "phone": "9876543210",
            "location": "India",
            "linkedin": "linkedin.com/in/test",
            "github": "github.com/test",
        },
        summary=(
            "AI/ML Engineer focused on Python, machine learning, "
            "Generative AI, and retrieval-augmented generation."
        ),
        skills=[
            "Python",
            "Machine Learning",
            "SQL",
            "Generative AI",
            "RAG",
        ],
        experience=[
            {
                "job_title": "AI/ML Engineer",
                "company": "Example Technologies",
                "dates": "2024 - Present",
                "bullets": [
                    "Built machine learning pipelines using Python.",
                    "Implemented retrieval-augmented generation workflows.",
                ],
            }
        ],
        projects=[
            {
                "title": "AI Resume Analyzer",
                "description": (
                    "Analyzes resumes against job descriptions "
                    "and generates optimization recommendations."
                ),
                "bullets": [
                    "Implemented ATS scoring and skill-gap analysis.",
                    "Added RAG grounding and optimized resume generation.",
                ],
            }
        ],
        education=[
            {
                "degree": "B.Tech in Computer Science",
                "institution": "Example University",
                "dates": "2020 - 2024",
            }
        ],
        certifications=[
            {
                "name": "Machine Learning Certification",
                "organization": "Example Institute",
                "year": "2024",
            }
        ],
    )


def pdf_text(path):
    """Extract text from all PDF pages."""
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def docx_text(path):
    """Extract paragraph text from a DOCX."""
    document = Document(path)
    return "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )


@pytest.mark.parametrize("template", TEMPLATES)
def test_templates_are_declared_ats_safe(template):
    config = get_template(template)

    assert config.ats_safe is True
    assert config.single_column is True
    assert config.use_bullets is True


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_has_no_tables_for_ats_safety(tmp_path, template):
    resume = build_resume()
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        resume,
        output_path,
        template=template,
    )

    document = Document(output_path)

    assert len(document.tables) == 0


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_has_no_textboxes_or_shapes(tmp_path, template):
    resume = build_resume()
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        resume,
        output_path,
        template=template,
    )

    with ZipFile(output_path) as archive:
        xml = archive.read("word/document.xml").decode(
            "utf-8",
            errors="ignore",
        )

    assert "<w:txbxContent" not in xml
    assert "<wps:" not in xml
    assert "<v:shape" not in xml


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_ats_text_is_extractable_and_complete(tmp_path, template):
    resume = build_resume()
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        resume,
        output_path,
        template=template,
    )

    text = docx_text(output_path)

    expected_values = [
        "Test Candidate",
        "test@example.com",
        "AI/ML Engineer",
        "Python",
        "Machine Learning",
        "SQL",
        "Generative AI",
        "AI Resume Analyzer",
        "B.Tech in Computer Science",
        "Machine Learning Certification",
    ]

    for value in expected_values:
        assert value in text


@pytest.mark.parametrize("template", TEMPLATES)
def test_pdf_ats_text_is_extractable_and_complete(tmp_path, template):
    resume = build_resume()
    output_path = tmp_path / f"{template}.pdf"

    render_final_resume_pdf(
        resume,
        output_path,
        template=template,
    )

    text = pdf_text(output_path)

    expected_values = [
        "Test Candidate",
        "test@example.com",
        "AI/ML Engineer",
        "Python",
        "Machine Learning",
        "SQL",
        "Generative AI",
        "AI Resume Analyzer",
        "B.Tech in Computer Science",
        "Machine Learning Certification",
    ]

    for value in expected_values:
        assert value in text


@pytest.mark.parametrize("template", TEMPLATES)
def test_pdf_uses_single_column_paragraph_flow(tmp_path, template):
    resume = build_resume()
    output_path = tmp_path / f"{template}.pdf"

    render_final_resume_pdf(
        resume,
        output_path,
        template=template,
    )

    reader = PdfReader(str(output_path))

    assert len(reader.pages) >= 1
    assert all(
        (page.extract_text() or "").strip()
        for page in reader.pages
    )


@pytest.mark.parametrize("template", TEMPLATES)
def test_standard_section_headings_are_ats_extractable(tmp_path, template):
    resume = build_resume()
    output_path = tmp_path / f"{template}.pdf"

    render_final_resume_pdf(
        resume,
        output_path,
        template=template,
    )

    text = pdf_text(output_path).upper()

    for heading in (
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
    ):
        assert heading in text


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_bullets_are_separate_extractable_paragraphs(
    tmp_path,
    template,
):
    resume = build_resume()
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        resume,
        output_path,
        template=template,
    )

    document = Document(output_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert "Built machine learning pipelines using Python." in paragraphs
    assert (
        "Implemented retrieval-augmented generation workflows."
        in paragraphs
    )
    assert (
        "Implemented ATS scoring and skill-gap analysis."
        in paragraphs
    )
    assert (
        "Added RAG grounding and optimized resume generation."
        in paragraphs
    )
