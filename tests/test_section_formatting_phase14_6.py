from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from app.final_resume_models import FinalResume
from app.resume_document_renderer import render_final_resume
from app.resume_pdf_renderer import render_final_resume_pdf
from app.templates.registry import get_template


def _resume_with_all_sections():
    return FinalResume(
        personal_info={
            "name": "Test Candidate",
            "email": "test@example.com",
            "phone": "1234567890",
        },
        summary="AI/ML engineer focused on Python and machine learning.",
        skills=["Python", "Machine Learning", "SQL"],
        experience=[
            {
                "job_title": "AI Engineer",
                "company": "Example Corp",
                "location": "Bengaluru",
                "dates": "2024 - 2026",
                "bullets": [
                    "Built machine learning pipelines.",
                    "Improved model evaluation workflows.",
                ],
            }
        ],
        projects=[
            {
                "title": "Resume Analyzer",
                "bullets": ["Built an ATS optimization pipeline."],
            }
        ],
        education=[
            {
                "degree": "B.Tech",
                "institution": "Example University",
                "year": "2024",
            }
        ],
        certifications=[
            {
                "name": "Machine Learning Certification",
                "organization": "Example Institute",
                "year": "2025",
            }
        ],
    )


def test_docx_section_heading_uses_template_spacing(tmp_path):
    resume = _resume_with_all_sections()
    output = tmp_path / "section_spacing.docx"

    render_final_resume(
        resume,
        output,
        template="modern",
    )

    document = Document(output)
    paragraphs = document.paragraphs

    heading = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text == "Experience"
    )

    config = get_template("modern")

    assert heading.paragraph_format.space_before.pt == config.section_space_before
    assert heading.paragraph_format.space_after.pt == config.section_space_after
    assert heading.paragraph_format.keep_with_next is True


def test_docx_item_spacing_uses_template_body_spacing(tmp_path):
    resume = _resume_with_all_sections()
    output = tmp_path / "item_spacing.docx"

    render_final_resume(
        resume,
        output,
        template="technical",
    )

    document = Document(output)
    paragraphs = document.paragraphs
    config = get_template("technical")

    secondary = next(
        paragraph
        for paragraph in paragraphs
        if "Example Corp" in paragraph.text
    )

    bullet = next(
        paragraph
        for paragraph in paragraphs
        if "Built machine learning pipelines." in paragraph.text
    )

    assert secondary.paragraph_format.space_after.pt == config.body_space_after
    assert bullet.paragraph_format.space_after.pt == config.body_space_after
    assert bullet.paragraph_format.keep_together is True


def test_docx_empty_sections_are_not_rendered(tmp_path):
    resume = FinalResume(
        personal_info={
            "name": "Test Candidate",
        },
        summary="Only summary content.",
        skills=[],
        experience=[],
        projects=[],
        education=[],
        certifications=[],
    )
    output = tmp_path / "empty_sections.docx"

    render_final_resume(resume, output)

    text = "\n".join(
        paragraph.text
        for paragraph in Document(output).paragraphs
    )

    assert "Professional Summary" in text
    assert "Skills" not in text
    assert "Experience" not in text
    assert "Projects" not in text
    assert "Education" not in text
    assert "Certifications" not in text


def test_pdf_styles_use_template_section_and_body_spacing():
    from app.resume_pdf_renderer import _build_styles

    config = get_template("modern")
    styles = _build_styles(config)

    assert styles["heading"].spaceBefore == config.section_space_before
    assert styles["heading"].spaceAfter == config.section_space_after
    assert styles["body"].spaceAfter == config.body_space_after
    assert styles["title"].spaceBefore == config.section_space_after
    assert styles["secondary"].spaceAfter == config.body_space_after
    assert styles["bullet"].spaceAfter == config.body_space_after


def test_pdf_all_sections_render_with_configured_formatting(tmp_path):
    resume = _resume_with_all_sections()
    output = tmp_path / "all_sections.pdf"

    render_final_resume_pdf(
        resume,
        output,
        template="ats_clean",
    )

    assert output.exists()
    assert output.stat().st_size > 0


def test_section_order_is_preserved_across_rendering(tmp_path):
    resume = _resume_with_all_sections()
    resume.section_order = [
        "projects",
        "summary",
        "skills",
        "experience",
        "education",
        "certifications",
    ]

    docx_output = tmp_path / "ordered.docx"
    pdf_output = tmp_path / "ordered.pdf"

    render_final_resume(
        resume,
        docx_output,
        template="modern",
    )
    render_final_resume_pdf(
        resume,
        pdf_output,
        template="modern",
    )

    document = Document(docx_output)
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text in {
            "Professional Summary",
            "Skills",
            "Experience",
            "Projects",
            "Education",
            "Certifications",
        }
    ]

    assert headings == [
        "Projects",
        "Professional Summary",
        "Skills",
        "Experience",
        "Education",
        "Certifications",
    ]
