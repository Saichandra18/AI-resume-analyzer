from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app import create_app


@pytest.fixture
def client():
    """Create a Flask test client."""
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def test_generation_e2e_produces_valid_resume_documents(
    client,
    monkeypatch,
    tmp_path,
):
    """Verify the complete generation route produces usable documents."""

    generated_dir = tmp_path / "generated_resumes"
    generated_dir.mkdir()

    monkeypatch.setattr(
        "app.routes.generation_routes.GENERATED_RESUMES_DIR",
        generated_dir,
    )

    class FakeFinalResume:
        """Minimal FinalResume-compatible object."""

        def to_dict(self):
            return {
                "personal_info": {
                    "name": "Test User",
                    "email": "test@example.com",
                },
                "summary": (
                    "AI/ML Engineer with Python and "
                    "Generative AI experience."
                ),
                "skills": [
                    "Python",
                    "Machine Learning",
                    "Generative AI",
                    "RAG",
                ],
                "experience": [],
                "projects": [
                    {
                        "title": "AI Resume Analyzer",
                        "description": (
                            "AI-powered resume optimization "
                            "and ATS analysis."
                        ),
                    }
                ],
                "education": [],
                "certifications": [],
                "section_order": [
                    "summary",
                    "skills",
                    "experience",
                    "projects",
                    "education",
                    "certifications",
                ],
                "metadata": {
                    "target_pages": 1,
                },
            }

    final_resume = FakeFinalResume()

    def fake_generate(**kwargs):
        """Generate real test DOCX and PDF artifacts."""

        docx_path = generated_dir / "optimized_resume.docx"
        pdf_path = generated_dir / "optimized_resume.pdf"

        # Create a real DOCX.
        document = Document()

        document.add_heading(
            "Test User",
            level=1,
        )

        document.add_paragraph(
            "AI/ML Engineer with Python and "
            "Generative AI experience."
        )

        document.add_heading(
            "Skills",
            level=2,
        )

        document.add_paragraph(
            "Python | Machine Learning | "
            "Generative AI | RAG"
        )

        document.add_heading(
            "Projects",
            level=2,
        )

        document.add_paragraph(
            "AI Resume Analyzer"
        )

        document.save(docx_path)

        # Create a real one-page PDF.
        from reportlab.pdfgen import canvas

        pdf = canvas.Canvas(str(pdf_path))

        pdf.drawString(
            72,
            750,
            "Test User",
        )

        pdf.drawString(
            72,
            720,
            "AI/ML Engineer with Python and "
            "Generative AI experience.",
        )

        pdf.drawString(
            72,
            690,
            "Skills: Python, Machine Learning, "
            "Generative AI, RAG",
        )

        pdf.drawString(
            72,
            660,
            "Projects: AI Resume Analyzer",
        )

        pdf.save()

        return {
            "final_resume": final_resume,
            "tailoring_plan": {},
            "rewrite_result": {},
            "length_control": {},
            "validation": {
                "valid": True,
                "errors": [],
                "warnings": [],
            },
            "document_path": docx_path,
            "pdf_path": pdf_path,
            "document_validation": {
                "valid": True,
                "target_pages": 1,
                "actual_pages": 1,
                "page_count_match": True,
            },
        }

    monkeypatch.setattr(
        "app.routes.generation_routes.generate_optimized_resume_documents",
        fake_generate,
    )

    response = client.post(
        "/api/resume/generate",
        json={
            "resume_text": (
                "Python AI ML engineer "
                "with Generative AI experience."
            ),
            "job_description": (
                "AI/ML Engineer with Python, "
                "Machine Learning, NLP, LLM and RAG."
            ),
            "target_pages": 1,
            "experience_level": "fresher",
            "target_job_title": "AI/ML Engineer",
        },
    )

    assert response.status_code == 200

    data = response.get_json()

    assert data["status"] == "success"

    assert data["document_path"].endswith(
        "optimized_resume.docx"
    )

    assert data["pdf_path"].endswith(
        "optimized_resume.pdf"
    )

    assert data["final_resume"]["summary"] == (
        "AI/ML Engineer with Python and "
        "Generative AI experience."
    )

    assert data["validation"]["valid"] is True

    # ---------------------------------------------------------
    # Verify DOCX is a real readable document.
    # ---------------------------------------------------------

    docx_path = generated_dir / "optimized_resume.docx"

    assert docx_path.exists()
    assert docx_path.stat().st_size > 0

    document = Document(str(docx_path))

    docx_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Test User" in docx_text
    assert "AI/ML Engineer" in docx_text
    assert "Python" in docx_text
    assert "AI Resume Analyzer" in docx_text

    # ---------------------------------------------------------
    # Verify PDF is a real readable document.
    # ---------------------------------------------------------

    pdf_path = generated_dir / "optimized_resume.pdf"

    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 0

    reader = PdfReader(str(pdf_path))

    assert len(reader.pages) == 1

    pdf_text = "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )

    assert "Test User" in pdf_text
    assert "AI/ML Engineer" in pdf_text
    assert "Python" in pdf_text
    assert "AI Resume Analyzer" in pdf_text


def test_generation_e2e_downloads_generated_docx_and_pdf(
    client,
    monkeypatch,
    tmp_path,
):
    """Verify generated DOCX and PDF can be downloaded."""

    generated_dir = tmp_path / "generated_resumes"
    generated_dir.mkdir()

    monkeypatch.setattr(
        "app.routes.generation_routes.GENERATED_RESUMES_DIR",
        generated_dir,
    )

    docx_path = generated_dir / "optimized_resume.docx"
    pdf_path = generated_dir / "optimized_resume.pdf"

    docx_path.write_bytes(
        b"generated-docx-content"
    )

    pdf_path.write_bytes(
        b"generated-pdf-content"
    )

    docx_response = client.get(
        "/api/resume/download"
        "?filename=optimized_resume.docx"
    )

    assert docx_response.status_code == 200
    assert docx_response.data == (
        b"generated-docx-content"
    )

    pdf_response = client.get(
        "/api/resume/download"
        "?filename=optimized_resume.pdf"
    )

    assert pdf_response.status_code == 200
    assert pdf_response.data == (
        b"generated-pdf-content"
    )

    assert pdf_response.headers[
        "Content-Type"
    ].startswith("application/pdf")