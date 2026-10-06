from pathlib import Path

from reportlab.pdfgen import canvas
from docx import Document

from app.final_ats_validation import compare_ats_results

from app.final_resume_models import FinalResume
from app.resume_generation_service import (
    generate_optimized_resume_docx,
    generate_optimized_resume_pdf,
    generate_optimized_resume_documents,
)


def _create_test_pdf(output_path):
    """Create a minimal valid one-page PDF for service tests."""
    pdf = canvas.Canvas(str(output_path))
    pdf.drawString(72, 720, "Test Resume")
    pdf.save()


def test_generate_optimized_resume_docx_includes_ats_evaluation(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI/ML Engineer",
        skills=[
            "Python",
            "Machine Learning",
        ],
        experience=[],
        projects=[],
        education=[],
        certifications=[],
    )

    generation_result = {
        "final_resume": final_resume,
        "validation": {
            "valid": True,
        },
    }

    expected_evaluation = {
        "before_score": 50,
        "after_score": 80,
        "improvement": 30,
        "before": {
            "skills": 50,
            "experience": None,
            "qualifications": None,
            "keywords": 50,
        },
        "after": {
            "skills": 80,
            "experience": None,
            "qualifications": None,
            "keywords": 80,
        },
    }

    def fake_generate_optimized_resume(**kwargs):
        return generation_result

    def fake_render_final_resume(**kwargs):
        output_path = Path(kwargs["output_path"])
        output_path.write_bytes(b"docx")
        return output_path

    def fake_evaluate_resume_optimization(**kwargs):
        assert kwargs["original_resume_data"] == {
            "skills": ["Python"],
        }

        assert kwargs["optimized_resume"] is final_resume
        assert kwargs["job_description"] == (
            "Python machine learning engineer"
        )

        return expected_evaluation

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        fake_generate_optimized_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume",
        fake_render_final_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        fake_evaluate_resume_optimization,
    )

    result = generate_optimized_resume_docx(
        resume_text="Original resume",
        output_path=tmp_path / "optimized.docx",
        resume_data={
            "skills": ["Python"],
        },
        job_description=(
            "Python machine learning engineer"
        ),
    )

    assert result["ats_evaluation"] == expected_evaluation
    assert result["document_path"].exists()


def test_generate_optimized_resume_docx_uses_empty_resume_data_when_missing(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI Engineer",
    )

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        lambda **kwargs: {
            "final_resume": final_resume,
            "validation": {
                "valid": True,
            },
        },
    )

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume",
        lambda **kwargs: kwargs["output_path"],
    )

    captured = {}

    def fake_evaluate(**kwargs):
        captured.update(kwargs)

        return {
            "before_score": 0,
            "after_score": 0,
            "improvement": 0,
            "before": {},
            "after": {},
        }

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        fake_evaluate,
    )

    result = generate_optimized_resume_docx(
        resume_text="Original resume",
        output_path=tmp_path / "optimized.docx",
        job_description="AI Engineer",
    )

    assert result["ats_evaluation"]["improvement"] == 0
    assert captured["original_resume_data"] == {}
    assert captured["optimized_resume"] is final_resume
    assert captured["job_description"] == "AI Engineer"


def test_generate_optimized_resume_pdf_includes_ats_evaluation(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI Engineer",
        skills=["Python", "Machine Learning"],
    )

    generation_result = {
        "final_resume": final_resume,
        "validation": {
            "valid": True,
        },
    }

    expected_evaluation = {
        "before_score": 40,
        "after_score": 75,
        "improvement": 35,
        "before": {},
        "after": {},
    }

    def fake_generate_optimized_resume(**kwargs):
        return generation_result

    def fake_render_final_resume_pdf(**kwargs):
        output_path = Path(kwargs["output_path"])

        _create_test_pdf(output_path)

        return output_path

    def fake_evaluate_resume_optimization(**kwargs):
        assert kwargs["optimized_resume"] is final_resume
        assert kwargs["job_description"] == "AI Engineer"

        return expected_evaluation

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        fake_generate_optimized_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume_pdf",
        fake_render_final_resume_pdf,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        fake_evaluate_resume_optimization,
    )

    result = generate_optimized_resume_pdf(
        resume_text="Original resume",
        output_path=tmp_path / "optimized.pdf",
        resume_data={"skills": ["Python"]},
        job_description="AI Engineer",
    )

    assert result["ats_evaluation"] == expected_evaluation

    # PDF-specific service result.
    assert result["pdf_path"].exists()


def test_generate_optimized_resume_documents_generates_docx_and_pdf(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI/ML Engineer",
        skills=["Python", "Machine Learning"],
    )

    generation_result = {
        "final_resume": final_resume,
        "validation": {
            "valid": True,
        },
    }

    calls = {
        "generation": 0,
        "docx": 0,
        "pdf": 0,
        "evaluation": 0,
    }

    ats_comparison = compare_ats_results(
        {
            "ats_score": 50,
            "skills": {"matched": ["Python"], "missing": []},
            "keywords": {"matched": ["Machine Learning"], "missing": []},
            "component_scores": {
                "skills": 50,
                "experience": 50,
                "qualifications": 50,
                "keywords": 50,
            },
        },
        {
            "ats_score": 85,
            "skills": {"matched": ["Python"], "missing": []},
            "keywords": {"matched": ["Machine Learning"], "missing": []},
            "component_scores": {
                "skills": 85,
                "experience": 85,
                "qualifications": 85,
                "keywords": 85,
            },
        },
    )

    expected_evaluation = {
        "before_score": 50,
        "after_score": 85,
        "improvement": 35,
        "before": {},
        "after": {},
        "comparison": ats_comparison,
    }

    def fake_generate_optimized_resume(**kwargs):
        calls["generation"] += 1

        return generation_result

    def fake_render_final_resume(**kwargs):
        calls["docx"] += 1

        output_path = Path(kwargs["output_path"])
        document = Document()

        for heading in (
            "SUMMARY",
            "SKILLS",
            "EXPERIENCE",
            "PROJECTS",
            "EDUCATION",
            "CERTIFICATIONS",
        ):
            document.add_heading(heading, level=1)

        document.add_paragraph("AI/ML Engineer")
        document.add_paragraph("Python Machine Learning")
        document.save(output_path)

        assert kwargs["resume"] is final_resume

        return output_path

    def fake_render_final_resume_pdf(**kwargs):
        calls["pdf"] += 1

        output_path = Path(kwargs["output_path"])

        # Create a real one-page PDF containing the standard
        # resume headings required by final document validation.
        pdf = canvas.Canvas(str(output_path))

        for index, line in enumerate(
            (
                "SUMMARY",
                "SKILLS",
                "EXPERIENCE",
                "PROJECTS",
                "EDUCATION",
                "CERTIFICATIONS",
                "AI/ML Engineer",
                "Python Machine Learning",
            )
        ):
            pdf.drawString(72, 750 - (20 * index), line)

        pdf.save()

        assert kwargs["resume"] is final_resume

        return output_path

    def fake_evaluate_resume_optimization(**kwargs):
        calls["evaluation"] += 1

        assert kwargs["optimized_resume"] is final_resume
        assert kwargs["job_description"] == "AI Engineer"

        return expected_evaluation

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        fake_generate_optimized_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume",
        fake_render_final_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume_pdf",
        fake_render_final_resume_pdf,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        fake_evaluate_resume_optimization,
    )

    result = generate_optimized_resume_documents(
        resume_text="Original resume",
        docx_output_path=tmp_path / "optimized.docx",
        pdf_output_path=tmp_path / "optimized.pdf",
        resume_data={"skills": ["Python"]},
        job_description="AI Engineer",
    )

    # Critical architectural guarantee:
    # optimization/generation runs only once for both formats.
    assert calls["generation"] == 1

    # Both renderers receive the same FinalResume.
    assert calls["docx"] == 1
    assert calls["pdf"] == 1

    # ATS evaluation also happens only once.
    assert calls["evaluation"] == 1

    assert result["document_path"].exists()
    assert result["pdf_path"].exists()
    assert result["ats_evaluation"] == expected_evaluation
    assert result["final_resume"] is final_resume

    # Phase 15.8:
    # Validate ATS safety plus the final DOCX and PDF artifacts.
    assert result["document_validation"]["valid"] is True
    assert result["document_validation"]["ats_safety"].passed is True
    assert result["document_validation"]["docx"]["valid"] is True
    assert result["document_validation"]["pdf"]["valid"] is True
    assert result["document_validation"]["pdf"]["page_count"] == 1
    assert result["document_validation"]["pdf"]["page_count_valid"] is True

def test_generate_optimized_resume_docx_forwards_template(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI Engineer",
        skills=["Python"],
    )

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        lambda **kwargs: {
            "final_resume": final_resume,
            "validation": {"valid": True},
        },
    )

    captured = {}

    def fake_render_final_resume(**kwargs):
        captured.update(kwargs)

        output_path = Path(kwargs["output_path"])
        output_path.write_bytes(b"docx")

        return output_path

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume",
        fake_render_final_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        lambda **kwargs: {
            "before_score": 0,
            "after_score": 0,
            "improvement": 0,
            "before": {},
            "after": {},
        },
    )

    result = generate_optimized_resume_docx(
        resume_text="Original resume",
        output_path=tmp_path / "optimized.docx",
        job_description="AI Engineer",
        template="modern",
    )

    assert captured["template"] == "modern"
    assert captured["resume"] is final_resume
    assert result["document_path"].exists()


def test_generate_optimized_resume_pdf_forwards_template(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI Engineer",
        skills=["Python"],
    )

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        lambda **kwargs: {
            "final_resume": final_resume,
            "validation": {"valid": True},
        },
    )

    captured = {}

    def fake_render_final_resume_pdf(**kwargs):
        captured.update(kwargs)

        output_path = Path(kwargs["output_path"])
        _create_test_pdf(output_path)

        return output_path

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume_pdf",
        fake_render_final_resume_pdf,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        lambda **kwargs: {
            "before_score": 0,
            "after_score": 0,
            "improvement": 0,
            "before": {},
            "after": {},
        },
    )

    result = generate_optimized_resume_pdf(
        resume_text="Original resume",
        output_path=tmp_path / "optimized.pdf",
        job_description="AI Engineer",
        template="technical",
    )

    assert captured["template"] == "technical"
    assert captured["resume"] is final_resume
    assert result["pdf_path"].exists()


def test_generate_optimized_resume_documents_forwards_template_to_both_renderers(
    monkeypatch,
    tmp_path,
):
    final_resume = FinalResume(
        summary="AI/ML Engineer",
        skills=["Python", "Machine Learning"],
    )

    monkeypatch.setattr(
        "app.resume_generation_service.generate_optimized_resume",
        lambda **kwargs: {
            "final_resume": final_resume,
            "validation": {"valid": True},
        },
    )

    captured = {
        "docx_template": None,
        "pdf_template": None,
    }

    def fake_render_final_resume(**kwargs):
        captured["docx_template"] = kwargs["template"]

        output_path = Path(kwargs["output_path"])
        output_path.write_bytes(b"docx")

        assert kwargs["resume"] is final_resume

        return output_path

    def fake_render_final_resume_pdf(**kwargs):
        captured["pdf_template"] = kwargs["template"]

        output_path = Path(kwargs["output_path"])
        _create_test_pdf(output_path)

        assert kwargs["resume"] is final_resume

        return output_path

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume",
        fake_render_final_resume,
    )

    monkeypatch.setattr(
        "app.resume_generation_service.render_final_resume_pdf",
        fake_render_final_resume_pdf,
    )

    ats_comparison = compare_ats_results(
        {
            "ats_score": 0,
            "skills": {"matched": [], "missing": []},
            "keywords": {"matched": [], "missing": []},
            "component_scores": {
                "skills": 0,
                "experience": 0,
                "qualifications": 0,
                "keywords": 0,
            },
        },
        {
            "ats_score": 0,
            "skills": {"matched": [], "missing": []},
            "keywords": {"matched": [], "missing": []},
            "component_scores": {
                "skills": 0,
                "experience": 0,
                "qualifications": 0,
                "keywords": 0,
            },
        },
    )

    monkeypatch.setattr(
        "app.resume_generation_service.evaluate_resume_optimization",
        lambda **kwargs: {
            "before_score": 0,
            "after_score": 0,
            "improvement": 0,
            "before": {},
            "after": {},
            "comparison": ats_comparison,
        },
    )

    result = generate_optimized_resume_documents(
        resume_text="Original resume",
        docx_output_path=tmp_path / "optimized.docx",
        pdf_output_path=tmp_path / "optimized.pdf",
        job_description="AI Engineer",
        template="ats_clean",
    )

    assert captured["docx_template"] == "ats_clean"
    assert captured["pdf_template"] == "ats_clean"

    assert result["document_path"].exists()
    assert result["pdf_path"].exists()