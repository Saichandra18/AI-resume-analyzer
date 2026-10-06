from copy import deepcopy

import pytest
from docx import Document
from pypdf import PdfReader

from app.final_resume_models import FinalResume
from app.resume_generation_service import (
    generate_optimized_resume_docx,
    generate_optimized_resume_pdf,
    generate_optimized_resume_documents,
)

from app.templates.registry import get_template


TEMPLATES = [
    "default",
    "ats_clean",
    "modern",
    "technical",
]


SAMPLE_RESUME_TEXT = """
John Doe
john.doe@example.com | +91 9876543210 | Bengaluru, India

PROFESSIONAL SUMMARY
AI/ML Engineer with experience building machine learning,
generative AI, and RAG-based applications using Python and SQL.

SKILLS
Python, SQL, Machine Learning, Generative AI, RAG, Flask, Pandas

EXPERIENCE
Data Analyst
ABC Technologies
2021 - Present
Analyzed business data using Python and SQL.
Built machine learning and AI-powered analytics solutions.

PROJECTS
AI Resume Analyzer
Built an AI-powered platform for resume analysis and optimization.

EDUCATION
Bachelor of Technology in Computer Science
University of Technology
2020 - 2024
"""


SAMPLE_RESUME_DATA = {
    "personal_info": {
        "name": "John Doe",
        "email": "john.doe@example.com",
        "phone": "+91 9876543210",
        "location": "Bengaluru, India",
    },
    "summary": (
        "AI/ML Engineer with experience building machine learning, "
        "generative AI, and RAG-based solutions."
    ),
    "skills": [
        "Python",
        "SQL",
        "Machine Learning",
        "Generative AI",
        "RAG",
        "Flask",
        "Pandas",
    ],
    "experience": {
        "section": [
            "Data Analyst",
            "ABC Technologies",
            "2021 - Present",
        ],
        "job_titles": ["Data Analyst"],
        "years": ["2021 - Present"],
    },
    "projects": [
        "AI Resume Analyzer",
        "Built an AI-powered platform for resume analysis and optimization.",
    ],
    "education": {
        "degrees": [
            "Bachelor of Technology in Computer Science",
        ],
        "years": ["2020", "2024"],
    },
}


SAMPLE_JOB_DESCRIPTION = """
We are looking for an AI/ML Engineer with skills in Python, SQL,
Machine Learning, Generative AI, RAG, Flask, and data analytics.
"""


def _normalize(text):
    return " ".join((text or "").lower().split())


def _read_docx(path):
    document = Document(str(path))
    return "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )


def _read_pdf(path):
    reader = PdfReader(str(path))
    return "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )


def _assert_core_content(text):
    normalized = _normalize(text)

    assert "john doe" in normalized
    assert "python" in normalized
    assert "machine learning" in normalized
    assert "generative ai" in normalized
    assert "resume" in normalized

    assert (
        "optimization" in normalized
        or "optimizing" in normalized
        or "optimiz" in normalized
    )


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_docx_generation_for_all_templates(
    tmp_path,
    template,
):
    output_path = tmp_path / f"e2e_{template}.docx"

    result = generate_optimized_resume_docx(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    assert output_path.exists()
    assert output_path.stat().st_size > 0

    document = Document(str(output_path))

    assert len(document.paragraphs) > 0
    assert result is not None

    _assert_core_content(_read_docx(output_path))


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_pdf_generation_for_all_templates(
    tmp_path,
    template,
):
    output_path = tmp_path / f"e2e_{template}.pdf"

    result = generate_optimized_resume_pdf(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    assert output_path.exists()
    assert output_path.stat().st_size > 0

    reader = PdfReader(str(output_path))

    assert len(reader.pages) >= 1
    assert result is not None

    _assert_core_content(_read_pdf(output_path))


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_combined_generation(
    tmp_path,
    template,
):
    docx_path = tmp_path / f"e2e_{template}.docx"
    pdf_path = tmp_path / f"e2e_{template}.pdf"

    result = generate_optimized_resume_documents(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        docx_output_path=str(docx_path),
        pdf_output_path=str(pdf_path),
        template=template,
    )

    assert result is not None

    assert docx_path.exists()
    assert docx_path.stat().st_size > 0

    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 0

    _assert_core_content(_read_docx(docx_path))
    _assert_core_content(_read_pdf(pdf_path))

    
@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_docx_is_single_column_and_has_no_tables(
    tmp_path,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    generate_optimized_resume_docx(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    document = Document(str(output_path))

    assert len(document.tables) == 0

    config = get_template(template)

    assert config.single_column is True
    assert config.ats_safe is True


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_pdf_is_reasonably_paginated(
    tmp_path,
    template,
):
    output_path = tmp_path / f"{template}.pdf"

    generate_optimized_resume_pdf(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    reader = PdfReader(str(output_path))

    assert 1 <= len(reader.pages) <= 3


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_docx_has_extractable_text(
    tmp_path,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    generate_optimized_resume_docx(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    text = _read_docx(output_path)

    assert len(text.strip()) > 100
    _assert_core_content(text)


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_pdf_has_extractable_text(
    tmp_path,
    template,
):
    output_path = tmp_path / f"{template}.pdf"

    generate_optimized_resume_pdf(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    text = _read_pdf(output_path)

    assert len(text.strip()) > 100
    _assert_core_content(text)


def test_e2e_template_selection_produces_valid_documents(
    tmp_path,
):
    outputs = {}

    for template in TEMPLATES:
        output_path = tmp_path / f"{template}.docx"

        generate_optimized_resume_docx(
            resume_text=SAMPLE_RESUME_TEXT,
            resume_data=deepcopy(SAMPLE_RESUME_DATA),
            job_description=SAMPLE_JOB_DESCRIPTION,
            output_path=str(output_path),
            template=template,
        )

        outputs[template] = _normalize(
            _read_docx(output_path)
        )

    for text in outputs.values():
        assert "john doe" in text
        assert "python" in text
        assert "machine learning" in text
        assert "resume" in text


def test_e2e_template_profiles_are_distinct():
    configs = [
        get_template(template)
        for template in TEMPLATES
    ]

    profiles = {
        (
            config.font_name,
            config.body_font_size,
            config.heading_font_size,
            config.name_font_size,
            config.top_margin,
            config.left_margin,
            config.name_alignment,
        )
        for config in configs
    }

    assert len(profiles) >= 3


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_generation_does_not_mutate_resume_data(
    tmp_path,
    template,
):
    resume_data = deepcopy(SAMPLE_RESUME_DATA)
    original = deepcopy(resume_data)

    output_path = tmp_path / f"{template}.docx"

    generate_optimized_resume_docx(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=resume_data,
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    assert resume_data == original


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_generated_docx_has_no_empty_document(
    tmp_path,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    generate_optimized_resume_docx(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    document = Document(str(output_path))

    non_empty = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    assert len(non_empty) >= 5


@pytest.mark.parametrize("template", TEMPLATES)
def test_e2e_generated_pdf_has_no_blank_document(
    tmp_path,
    template,
):
    output_path = tmp_path / f"{template}.pdf"

    generate_optimized_resume_pdf(
        resume_text=SAMPLE_RESUME_TEXT,
        resume_data=deepcopy(SAMPLE_RESUME_DATA),
        job_description=SAMPLE_JOB_DESCRIPTION,
        output_path=str(output_path),
        template=template,
    )

    reader = PdfReader(str(output_path))

    extracted = "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )

    assert extracted.strip()


def test_e2e_default_template_remains_available():
    config = get_template("default")

    assert config is not None
    assert config.name == "default"


@pytest.mark.parametrize(
    "template",
    ["default", "ats_clean", "modern", "technical"],
)
def test_e2e_template_is_ats_safe(template):
    config = get_template(template)

    assert config.ats_safe is True
    assert config.single_column is True