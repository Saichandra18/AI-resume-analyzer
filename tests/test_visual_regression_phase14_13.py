from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from app.resume_document_renderer import render_final_resume
from app.resume_pdf_renderer import render_final_resume_pdf
from app.templates.registry import get_template


TEMPLATES = ["default", "ats_clean", "modern", "technical"]


@pytest.fixture
def sample_resume():
    from app.final_resume_models import FinalResume

    return FinalResume(
        personal_info={
            "name": "John Doe",
            "email": "john.doe@example.com",
            "phone": "+91 9876543210",
            "location": "Bengaluru, India",
        },
        summary=(
            "AI/ML Engineer specializing in machine learning, "
            "generative AI, and RAG applications."
        ),
        skills=[
            "Python",
            "SQL",
            "Machine Learning",
            "Generative AI",
            "RAG",
            "Flask",
        ],
        experience=[
            {
                "title": "AI/ML Engineer",
                "company": "Tech Company",
                "dates": "2021 - Present",
                "bullets": [
                    "Built machine learning applications using Python.",
                    "Developed RAG-based AI solutions.",
                ],
            }
        ],
        projects=[
            {
                "name": "AI Resume Analyzer",
                "description": "Built an AI-powered resume optimization platform.",
                "technologies": ["Python", "Flask", "RAG"],
            }
        ],
        education=[
            {
                "degree": "Bachelor of Technology in Computer Science",
                "institution": "University of Technology",
                "dates": "2020 - 2024",
            }
        ],
        certifications=[],
    )


@pytest.mark.parametrize("template", TEMPLATES)
def test_all_templates_generate_valid_docx(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(output_path),
        template=template,
    )

    assert output_path.exists()
    assert output_path.stat().st_size > 0

    document = Document(str(output_path))

    assert len(document.paragraphs) > 0


@pytest.mark.parametrize("template", TEMPLATES)
def test_all_templates_generate_valid_pdf(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.pdf"

    render_final_resume_pdf(
        sample_resume,
        str(output_path),
        template=template,
    )

    assert output_path.exists()
    assert output_path.stat().st_size > 0

    reader = PdfReader(str(output_path))

    assert len(reader.pages) >= 1


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_name_format_matches_template(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(output_path),
        template=template,
    )

    document = Document(str(output_path))

    name_paragraph = document.paragraphs[0]

    config = get_template(template)

    assert name_paragraph.text == "John Doe"

    run = name_paragraph.runs[0]

    assert run.bold == config.name_bold
    assert run.font.size is not None
    assert run.font.size.pt == pytest.approx(config.name_font_size)


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_alignment_matches_template(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(output_path),
        template=template,
    )

    document = Document(str(output_path))

    name_paragraph = document.paragraphs[0]

    config = get_template(template)

    alignment_map = {
        "left": 0,
        "center": 1,
        "right": 2,
        "justify": 3,
    }

    expected_alignment = alignment_map.get(
        config.name_alignment
    )

    if expected_alignment is not None:
        assert name_paragraph.alignment == expected_alignment


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_margins_match_template(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(output_path),
        template=template,
    )

    document = Document(str(output_path))
    section = document.sections[0]

    config = get_template(template)

    assert section.top_margin.inches == pytest.approx(
        config.top_margin
    )
    assert section.bottom_margin.inches == pytest.approx(
        config.bottom_margin
    )
    assert section.left_margin.inches == pytest.approx(
        config.left_margin
    )
    assert section.right_margin.inches == pytest.approx(
        config.right_margin
    )


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_contains_no_tables(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(output_path),
        template=template,
    )

    document = Document(str(output_path))

    assert len(document.tables) == 0


@pytest.mark.parametrize("template", TEMPLATES)
def test_docx_preserves_core_content(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(output_path),
        template=template,
    )

    document = Document(str(output_path))

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    normalized = " ".join(text.lower().split())

    assert "john doe" in normalized
    assert "python" in normalized
    assert "machine learning" in normalized
    assert "generative ai" in normalized
    assert "resume" in normalized
    assert "optimization" in normalized or "optimiz" in normalized


@pytest.mark.parametrize("template", TEMPLATES)
def test_pdf_preserves_core_content(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.pdf"

    render_final_resume_pdf(
        sample_resume,
        str(output_path),
        template=template,
    )

    reader = PdfReader(str(output_path))

    text = "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )

    normalized = " ".join(text.lower().split())

    assert "john doe" in normalized
    assert "python" in normalized
    assert "machine learning" in normalized
    assert "generative ai" in normalized
    assert "resume" in normalized
    assert (
        "optimization" in normalized
        or "optimizing" in normalized
        or "optimiz" in normalized
    )


def test_templates_have_distinct_formatting_profiles():
    configs = [get_template(template) for template in TEMPLATES]

    profiles = {
        (
            config.font_name,
            config.body_font_size,
            config.heading_font_size,
            config.name_font_size,
            config.top_margin,
            config.bottom_margin,
            config.left_margin,
            config.right_margin,
            config.name_alignment,
        )
        for config in configs
    }

    assert len(profiles) >= 3


@pytest.mark.parametrize("template", TEMPLATES)
def test_template_is_single_column_and_ats_safe(template):
    config = get_template(template)

    assert config.single_column is True
    assert config.ats_safe is True


@pytest.mark.parametrize("template", TEMPLATES)
def test_pdf_page_count_is_reasonable(
    tmp_path,
    sample_resume,
    template,
):
    output_path = tmp_path / f"{template}.pdf"

    render_final_resume_pdf(
        sample_resume,
        str(output_path),
        template=template,
    )

    reader = PdfReader(str(output_path))

    assert 1 <= len(reader.pages) <= 3


@pytest.mark.parametrize("template", TEMPLATES)
def test_rendering_does_not_mutate_resume(
    tmp_path,
    sample_resume,
    template,
):
    original = repr(sample_resume.__dict__)

    docx_path = tmp_path / f"{template}.docx"

    render_final_resume(
        sample_resume,
        str(docx_path),
        template=template,
    )

    assert repr(sample_resume.__dict__) == original