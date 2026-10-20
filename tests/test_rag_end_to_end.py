"""End-to-end tests for the real RAG knowledge-base workflow."""

from app.rag.knowledge.service import KnowledgeBaseService
from app.rag.config import RAGConfig
from app.rag.models import RetrievedChunk
from app.rag.optimization.service import RAGOptimizationService
from app.rag.models import DocumentChunk
from app.rag.optimization.service import RAGOptimizationService
from app.resume_generation_pipeline import generate_optimized_resume


def _jd_analysis():
    """Return a representative JD analysis."""

    return {
        "cleaned_text": (
            "Python machine learning engineer "
            "with SQL and Generative AI experience."
        ),
        "required_skills": [
            "Python",
            "Machine Learning",
            "SQL",
            "Generative AI",
        ],
        "qualifications": [],
        "experience": [],
        "keywords": [
            "Python",
            "machine learning",
            "SQL",
            "Generative AI",
        ],
    }


def test_real_knowledge_base_indexes_builtin_sources():
    """Built-in knowledge sources should be indexed successfully."""

    service = KnowledgeBaseService()

    assert service.is_indexed() is False

    result = service.index_all()

    assert result is not None
    assert service.is_indexed() is True
    assert service.indexed_count() > 0


def test_real_knowledge_base_reports_source_status():
    """All approved knowledge sources should have status information."""

    service = KnowledgeBaseService()

    service.index_all()

    status = service.source_status()

    assert status
    assert isinstance(status, (list, tuple, dict))


def test_real_knowledge_base_retrieves_builtin_knowledge():
    """Indexed built-in knowledge should be retrievable for a JD."""

    service = KnowledgeBaseService()

    service.index_all()

    results = service.retrieve(_jd_analysis())

    assert isinstance(results, list)
    assert results

    for result in results:
        assert result.chunk_id
        assert result.text
        assert result.document_id is not None
        assert result.score is not None


def test_real_knowledge_base_retrieval_returns_retrieved_chunks():
    """Real retrieval should return the expected RAG model."""

    from app.rag.models import RetrievedChunk

    service = KnowledgeBaseService()

    service.index_all()

    results = service.retrieve(_jd_analysis())

    assert results
    assert all(
        isinstance(result, RetrievedChunk)
        for result in results
    )


def test_real_knowledge_base_retrieval_contains_knowledge_metadata():
    """Retrieved knowledge should retain source metadata."""

    service = KnowledgeBaseService()

    service.index_all()

    results = service.retrieve(_jd_analysis())

    assert results

    assert all(
        isinstance(result.metadata, dict)
        for result in results
    )


def test_real_knowledge_base_index_count_matches_retrieval_state():
    """Indexing should make the knowledge base report indexed records."""

    service = KnowledgeBaseService()

    assert service.indexed_count() == 0

    service.index_all()

    indexed_count = service.indexed_count()

    assert indexed_count > 0
    assert service.is_indexed() is True


def test_real_knowledge_base_retrieve_with_metadata():
    """Metadata retrieval should expose both chunks and index state."""

    service = KnowledgeBaseService()

    service.index_all()

    result = service.retrieve_with_metadata(
        _jd_analysis()
    )

    assert isinstance(result, dict)

    assert "retrieved_chunks" in result
    assert "count" in result
    assert "indexed_count" in result
    assert "metadata" in result

    assert result["retrieved_chunks"]
    assert result["count"] == len(
        result["retrieved_chunks"]
    )

    assert result["indexed_count"] > 0

    assert result["metadata"]["knowledge_base"] is True
    assert (
        result["metadata"]["indexed_count"]
        == result["indexed_count"]
    )

class FakeOptimizationOrchestrator:
    """Provide deterministic candidate optimization results."""

    def __init__(self):
        self.calls = []

    def optimize(self, jd_analysis, resume_context=None):
        self.calls.append(
            {
                "jd_analysis": jd_analysis,
                "resume_context": resume_context,
            }
        )

        return {
            "retrieved_chunks": [],
            "evidence": [],
            "recommendations": [],
            "safety_results": [],
            "signals": None,
            "phase11_payload": {
                "skills_to_emphasize": ["Python"],
                "experience_to_emphasize": [],
                "keywords_to_emphasize": [
                    "machine learning"
                ],
                "emphasis_instructions": [],
                "grounding": [],
                "metadata": {},
            },
            "metadata": {},
        }


def _enabled_config():
    """Return an enabled RAG configuration."""

    config = RAGConfig()
    config.enabled = True
    return config


def _sample_resume_data():
    """Return representative candidate resume data."""

    return {
        "summary": (
            "Python developer with machine learning "
            "experience."
        ),
        "skills": [
            "Python",
            "SQL",
            "Machine Learning",
        ],
        "experience": [
            {
                "company": "Example Company",
                "role": "Python Developer",
                "description": (
                    "Built machine learning applications."
                ),
            }
        ],
        "projects": [
            {
                "name": "AI Resume Analyzer",
                "description": (
                    "Built an AI-powered resume "
                    "optimization system."
                ),
            }
        ],
        "education": ["B.Tech"],
    }


def test_real_knowledge_base_integrates_with_rag_service():
    """Real KB retrieval should reach the RAG service."""

    knowledge_base = KnowledgeBaseService()
    knowledge_base.index_all()

    orchestrator = FakeOptimizationOrchestrator()

    service = RAGOptimizationService(
        config=_enabled_config(),
        orchestrator=orchestrator,
        knowledge_base_service=knowledge_base,
    )

    result = service.optimize(
        _jd_analysis(),
        resume_data=_sample_resume_data(),
    )

    assert result["enabled"] is True

    knowledge_context = result["knowledge_context"]

    assert knowledge_context["enabled"] is True
    assert knowledge_context["count"] > 0
    assert knowledge_context["retrieved_chunks"]

    assert all(
        isinstance(chunk, RetrievedChunk)
        for chunk in knowledge_context[
            "retrieved_chunks"
        ]
    )


def test_real_knowledge_base_produces_guidance():
    """Real KB retrieval should produce adapted guidance."""

    knowledge_base = KnowledgeBaseService()
    knowledge_base.index_all()

    service = RAGOptimizationService(
        config=_enabled_config(),
        orchestrator=FakeOptimizationOrchestrator(),
        knowledge_base_service=knowledge_base,
    )

    result = service.optimize(
        _jd_analysis(),
        resume_data=_sample_resume_data(),
    )

    knowledge_context = result["knowledge_context"]

    assert knowledge_context["guidance"]

    for guidance in knowledge_context["guidance"]:
        assert guidance["text"]
        assert guidance["source_chunk_id"]
        assert guidance["document_id"] is not None
        assert guidance["category"]
        assert "metadata" in guidance


def test_knowledge_guidance_remains_separate_from_candidate_evidence():
    """Knowledge guidance must not become candidate evidence."""

    knowledge_base = KnowledgeBaseService()
    knowledge_base.index_all()

    service = RAGOptimizationService(
        config=_enabled_config(),
        orchestrator=FakeOptimizationOrchestrator(),
        knowledge_base_service=knowledge_base,
    )

    result = service.optimize(
        _jd_analysis(),
        resume_data=_sample_resume_data(),
    )

    knowledge_context = result["knowledge_context"]

    assert "guidance" in knowledge_context

    for guidance in knowledge_context["guidance"]:
        assert "evidence_id" not in guidance
        assert "source_chunk_ids" not in guidance


def test_knowledge_guidance_does_not_replace_phase11_payload():
    """Knowledge guidance should coexist with Phase 11 optimization."""

    knowledge_base = KnowledgeBaseService()
    knowledge_base.index_all()

    orchestrator = FakeOptimizationOrchestrator()

    service = RAGOptimizationService(
        config=_enabled_config(),
        orchestrator=orchestrator,
        knowledge_base_service=knowledge_base,
    )

    result = service.optimize(
        _jd_analysis(),
        resume_data=_sample_resume_data(),
    )

    payload = result["phase11_payload"]

    assert payload["skills_to_emphasize"] == [
        "Python"
    ]

    assert payload["keywords_to_emphasize"] == [
        "machine learning"
    ]

    assert result["knowledge_context"]["guidance"]


def test_real_knowledge_base_preserves_resume_aware_optimization():
    """Knowledge retrieval must coexist with resume-aware optimization."""

    knowledge_base = KnowledgeBaseService()
    knowledge_base.index_all()

    orchestrator = FakeOptimizationOrchestrator()

    resume_data = _sample_resume_data()

    service = RAGOptimizationService(
        config=_enabled_config(),
        orchestrator=orchestrator,
        knowledge_base_service=knowledge_base,
    )

    result = service.optimize(
        _jd_analysis(),
        resume_data=resume_data,
    )

    assert len(orchestrator.calls) == 1

    call = orchestrator.calls[0]

    assert call["resume_context"] is not None

    assert (
        call["resume_context"]["skills"]
        == resume_data["skills"]
    )

    assert result["metadata"][
        "resume_context_available"
    ] is True

    assert result["metadata"][
        "knowledge_base_enabled"
    ] is True


def test_real_knowledge_base_metadata_is_exposed():
    """Knowledge-base runtime metadata should be preserved."""

    knowledge_base = KnowledgeBaseService()
    knowledge_base.index_all()

    service = RAGOptimizationService(
        config=_enabled_config(),
        orchestrator=FakeOptimizationOrchestrator(),
        knowledge_base_service=knowledge_base,
    )

    result = service.optimize(
        _jd_analysis()
    )

    knowledge_context = result["knowledge_context"]

    assert knowledge_context["metadata"][
        "knowledge_base"
    ] is True

    assert (
        knowledge_context["metadata"]["indexed_count"]
        == knowledge_base.indexed_count()
    )

    assert result["metadata"][
        "knowledge_base_enabled"
    ] is True

def _sample_candidate_resume_data():
    """Return candidate data used by the real RAG pipeline."""

    return {
        "personal_info": {
            "name": "Test Candidate",
            "email": "candidate@example.com",
        },
        "summary": (
            "Python developer with machine learning experience."
        ),
        "skills": [
            "Python",
            "SQL",
            "Machine Learning",
        ],
        "experience": [
            {
                "job_title": "Python Developer",
                "company": "Example Company",
                "dates": "2024 - Present",
                "bullets": [
                    "Built Python machine learning applications."
                ],
            }
        ],
        "projects": [
            {
                "title": "AI Resume Analyzer",
                "bullets": [
                    "Built an AI-powered resume analyzer."
                ],
            }
        ],
        "education": [
            {
                "degree": "B.Tech",
                "institution": "Example University",
                "year": "2024",
            }
        ],
        "certifications": [],
    }


def _sample_candidate_chunks():
    """Return candidate-owned chunks for real RAG retrieval."""

    return [
        DocumentChunk(
            chunk_id="candidate-skills",
            document_id="candidate-resume",
            text=(
                "Candidate skills include Python, SQL, "
                "and Machine Learning."
            ),
            metadata={
                "section": "skills",
                "source": "candidate_resume",
            },
        ),
        DocumentChunk(
            chunk_id="candidate-experience",
            document_id="candidate-resume",
            text=(
                "Worked as a Python Developer building "
                "machine learning applications."
            ),
            metadata={
                "section": "experience",
                "source": "candidate_resume",
            },
        ),
        DocumentChunk(
            chunk_id="candidate-project",
            document_id="candidate-resume",
            text=(
                "Built an AI Resume Analyzer using "
                "Python and machine learning."
            ),
            metadata={
                "section": "projects",
                "source": "candidate_resume",
            },
        ),
    ]


def _candidate_jd_analysis():
    """Return a JD analysis matching the candidate evidence."""

    return {
        "cleaned_text": (
            "Python machine learning engineer "
            "with SQL experience."
        ),
        "required_skills": [
            "Python",
            "Machine Learning",
            "SQL",
        ],
        "qualifications": [],
        "experience": [],
        "keywords": [
            "Python",
            "machine learning",
            "SQL",
        ],
    }


def _deterministic_rewrite_result():
    """Return deterministic Phase 10 output."""

    return {
        "summary": {
            "original": (
                "Python developer with machine "
                "learning experience."
            ),
            "rewritten": (
                "Python developer with machine "
                "learning experience."
            ),
            "changes": [],
            "validation": {
                "meaning_preserved": True,
                "unsupported_claims": [],
                "grounding_status": "grounded",
            },
        },
        "experience": [
            {
                "original": (
                    "Built Python machine learning applications."
                ),
                "rewritten": (
                    "Built Python machine learning applications."
                ),
                "changes": [],
                "validation": {
                    "meaning_preserved": True,
                    "unsupported_claims": [],
                    "grounding_status": "grounded",
                },
            }
        ],
        "projects": [
            {
                "original": (
                    "Built an AI-powered resume analyzer."
                ),
                "rewritten": (
                    "Built an AI-powered resume analyzer."
                ),
                "changes": [],
                "validation": {
                    "meaning_preserved": True,
                    "unsupported_claims": [],
                    "grounding_status": "grounded",
                },
            }
        ],
        "skills": {
            "original": [
                "Python",
                "SQL",
                "Machine Learning",
            ],
            "rewritten": [
                "Python",
                "SQL",
                "Machine Learning",
            ],
            "changes": [],
            "validation": {
                "meaning_preserved": True,
                "unsupported_claims": [],
                "grounding_status": "grounded",
            },
        },
        "validation": {
            "valid": True,
            "unsupported_claims": [],
        },
    }


def test_real_rag_service_retrieves_candidate_evidence():
    """The real RAG service should retrieve candidate-owned chunks."""

    service = RAGOptimizationService(
        config=_enabled_config()
    )

    chunks = _sample_candidate_chunks()

    indexed_ids = service.index_chunks(chunks)

    assert indexed_ids == [
        "candidate-skills",
        "candidate-experience",
        "candidate-project",
    ]

    result = service.optimize(
        _candidate_jd_analysis(),
        resume_data=_sample_candidate_resume_data(),
    )

    assert result["enabled"] is True
    assert result["retrieved_chunks"]

    retrieved_ids = {
        chunk.chunk_id
        for chunk in result["retrieved_chunks"]
    }

    assert retrieved_ids.intersection(
        {
            "candidate-skills",
            "candidate-experience",
            "candidate-project",
        }
    )


def test_real_rag_service_produces_phase11_payload():
    """Real candidate RAG should produce a Phase 11 payload."""

    service = RAGOptimizationService(
        config=_enabled_config()
    )

    service.index_chunks(
        _sample_candidate_chunks()
    )

    result = service.optimize(
        _candidate_jd_analysis(),
        resume_data=_sample_candidate_resume_data(),
    )

    payload = result["phase11_payload"]

    assert isinstance(payload, dict)

    assert "skills_to_emphasize" in payload
    assert "experience_to_emphasize" in payload
    assert "keywords_to_emphasize" in payload
    assert "emphasis_instructions" in payload

    assert result["metadata"]["rag_enabled"] is True
    assert result["metadata"][
        "resume_context_available"
    ] is True


def test_real_rag_flows_into_final_resume(monkeypatch):
    """Real RAG should flow through Phase 11 into FinalResume."""

    monkeypatch.setattr(
        "app.resume_generation_pipeline.rewrite_resume",
        lambda **kwargs: _deterministic_rewrite_result(),
    )

    rag_service = RAGOptimizationService(
        config=_enabled_config()
    )

    rag_service.index_chunks(
        _sample_candidate_chunks()
    )

    resume_data = _sample_candidate_resume_data()

    result = generate_optimized_resume(
        resume_text=(
            "Python developer with machine learning "
            "experience."
        ),
        resume_data=resume_data,
        job_description=(
            "Python machine learning engineer "
            "with SQL experience."
        ),
        ats_result={
            "keywords": [
                "Python",
                "machine learning",
                "SQL",
            ],
            "skills": [
                "Python",
                "Machine Learning",
                "SQL",
            ],
        },
        rag_service=rag_service,
    )

    assert result["rag_optimization"] is not None

    rag_payload = result["rag_optimization"]

    assert isinstance(rag_payload, dict)

    assert "skills_to_emphasize" in rag_payload
    assert "experience_to_emphasize" in rag_payload
    assert "keywords_to_emphasize" in rag_payload

    final_resume = result["final_resume"]

    assert final_resume is not None
    assert result["validation"]["valid"] is True

    assert final_resume.personal_info["name"] == (
        "Test Candidate"
    )

    assert final_resume.summary.startswith(
        "Python developer with machine learning experience."
    )

    # Supported JD keywords now live in Skills, not in a raw keyword
    # sentence appended to the summary.
    assert "core areas include" not in final_resume.summary.lower()
    assert "sql" in [skill.lower() for skill in final_resume.skills]


def test_real_rag_payload_is_attached_to_tailoring_plan():
    """The real RAG payload should remain attached to Phase 9 context."""

    service = RAGOptimizationService(
        config=_enabled_config()
    )

    service.index_chunks(
        _sample_candidate_chunks()
    )

    result = generate_optimized_resume(
        resume_text=(
            "Python developer with machine learning "
            "experience."
        ),
        resume_data=_sample_candidate_resume_data(),
        job_description=(
            "Python machine learning engineer "
            "with SQL experience."
        ),
        ats_result={
            "keywords": [
                "Python",
                "machine learning",
                "SQL",
            ],
            "skills": [
                "Python",
                "Machine Learning",
                "SQL",
            ],
        },
        tailoring_plan={
            "tailoring_plan": {
                "priorities": [],
                "summary": [],
                "skills": [],
                "experience": [],
                "keywords": [],
                "qualifications": [],
                "warnings": [],
            }
        },
        rag_service=service,
    )

    tailoring_plan = result["tailoring_plan"]

    assert "tailoring_plan" in tailoring_plan

    nested_plan = tailoring_plan[
        "tailoring_plan"
    ]

    assert "rag_optimization" in nested_plan

    assert (
        nested_plan["rag_optimization"]
        == result["rag_optimization"]
    )


def test_real_rag_does_not_inject_knowledge_as_candidate_claim():
    """RAG must not fabricate candidate experience."""

    service = RAGOptimizationService(
        config=_enabled_config()
    )

    service.index_chunks(
        _sample_candidate_chunks()
    )

    result = service.optimize(
        _candidate_jd_analysis(),
        resume_data=_sample_candidate_resume_data(),
    )

    candidate_text = str(
        _sample_candidate_resume_data()
    )

    rag_payload_text = str(
        result["phase11_payload"]
    )

    # Candidate-owned concepts are allowed.
    assert "Python" in candidate_text
    assert "Machine Learning" in candidate_text

    # Generic unsupported technologies should not
    # appear merely because they could be relevant
    # to the job.
    assert "Kubernetes" not in rag_payload_text
    assert "Terraform" not in rag_payload_text
    assert "Java" not in rag_payload_text

def test_real_rag_final_resume_renders_to_docx(
    monkeypatch,
    tmp_path,
):
    """Real RAG output should render successfully to a DOCX."""

    from docx import Document

    from app.resume_generation_service import render_final_resume

    monkeypatch.setattr(
        "app.resume_generation_pipeline.rewrite_resume",
        lambda **kwargs: _deterministic_rewrite_result(),
    )

    rag_service = RAGOptimizationService(
        config=_enabled_config()
    )

    rag_service.index_chunks(
        _sample_candidate_chunks()
    )

    resume_data = _sample_candidate_resume_data()

    result = generate_optimized_resume(
        resume_text=(
            "Python developer with machine learning "
            "experience."
        ),
        resume_data=resume_data,
        job_description=(
            "Python machine learning engineer "
            "with SQL experience."
        ),
        ats_result={
            "keywords": [
                "Python",
                "machine learning",
                "SQL",
            ],
            "skills": [
                "Python",
                "Machine Learning",
                "SQL",
            ],
        },
        rag_service=rag_service,
    )

    assert result["rag_optimization"] is not None
    assert result["final_resume"] is not None

    output_path = tmp_path / "optimized_resume.docx"

    rendered_path = render_final_resume(
        resume=result["final_resume"],
        output_path=output_path,
    )

    assert rendered_path == output_path
    assert output_path.exists()
    assert output_path.is_file()
    assert output_path.stat().st_size > 0

    document = Document(output_path)

    document_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Test Candidate" in document_text
    assert "candidate@example.com" in document_text
    assert "Python" in document_text
    assert "Machine Learning" in document_text
    assert "SQL" in document_text


def test_real_rag_docx_does_not_contain_rag_metadata(
    monkeypatch,
    tmp_path,
):
    """RAG metadata must not leak into the generated resume."""

    from docx import Document

    from app.resume_generation_service import render_final_resume

    monkeypatch.setattr(
        "app.resume_generation_pipeline.rewrite_resume",
        lambda **kwargs: _deterministic_rewrite_result(),
    )

    rag_service = RAGOptimizationService(
        config=_enabled_config()
    )

    rag_service.index_chunks(
        _sample_candidate_chunks()
    )

    result = generate_optimized_resume(
        resume_text=(
            "Python developer with machine learning "
            "experience."
        ),
        resume_data=_sample_candidate_resume_data(),
        job_description=(
            "Python machine learning engineer "
            "with SQL experience."
        ),
        ats_result={
            "keywords": [
                "Python",
                "machine learning",
                "SQL",
            ],
            "skills": [
                "Python",
                "Machine Learning",
                "SQL",
            ],
        },
        rag_service=rag_service,
    )

    output_path = tmp_path / "optimized_resume.docx"

    render_final_resume(
        resume=result["final_resume"],
        output_path=output_path,
    )

    document = Document(output_path)

    document_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "rag_optimization" not in document_text
    assert "knowledge_context" not in document_text
    assert "source_chunk_id" not in document_text
    assert "evidence_id" not in document_text
    assert "candidate-skills" not in document_text
    assert "candidate-experience" not in document_text
    assert "candidate-project" not in document_text


def test_real_rag_docx_preserves_candidate_sections(
    monkeypatch,
    tmp_path,
):
    """Generated DOCX should preserve the candidate's core sections."""

    from docx import Document

    from app.resume_generation_service import render_final_resume

    monkeypatch.setattr(
        "app.resume_generation_pipeline.rewrite_resume",
        lambda **kwargs: _deterministic_rewrite_result(),
    )

    rag_service = RAGOptimizationService(
        config=_enabled_config()
    )

    rag_service.index_chunks(
        _sample_candidate_chunks()
    )

    result = generate_optimized_resume(
        resume_text=(
            "Python developer with machine learning "
            "experience."
        ),
        resume_data=_sample_candidate_resume_data(),
        job_description=(
            "Python machine learning engineer "
            "with SQL experience."
        ),
        ats_result={
            "keywords": [
                "Python",
                "machine learning",
                "SQL",
            ],
            "skills": [
                "Python",
                "Machine Learning",
                "SQL",
            ],
        },
        rag_service=rag_service,
    )

    output_path = tmp_path / "optimized_resume.docx"

    render_final_resume(
        resume=result["final_resume"],
        output_path=output_path,
    )

    document = Document(output_path)

    document_text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "SUMMARY" in document_text.upper()
    assert "SKILLS" in document_text.upper()
    assert "EXPERIENCE" in document_text.upper()
    assert "PROJECTS" in document_text.upper()
    assert "EDUCATION" in document_text.upper()